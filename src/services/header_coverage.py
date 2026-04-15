# -*- coding: utf-8 -*-
"""
表头覆盖检查：与导入推断（excel_to_ods._infer_field_columns）及 sheet_header_slugs 对齐，
找出「当前 YAML 未命中」的列名，并给出建议英文 slug（供人工写入 sheet_header_slugs 或跑 sync 工具）。
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

from config.field_mapping import get_field_mapping_for_sheet
from src.ingestion.excel_to_ods import _infer_field_columns, _normalize_colname
from src.services.field_slug import normalize_header_key, suggest_field_slug


def _norm(s: Any) -> str:
    return normalize_header_key(s)


def _detect_header_row(ws: Any, *, max_scan_rows: int = 12, min_nonempty: int = 5) -> list[str] | None:
    max_r = min(max_scan_rows, int(ws.max_row or 0))
    max_c = int(ws.max_column or 0)
    if max_r <= 0 or max_c <= 0:
        return None
    best: tuple[int, list[str]] | None = None
    for r in range(1, max_r + 1):
        texts: list[str] = []
        nonempty = 0
        for c in range(1, max_c + 1):
            v = ws.cell(row=r, column=c).value
            t = "" if v is None else str(v).strip()
            texts.append(t)
            if t:
                nonempty += 1
        if nonempty >= min_nonempty:
            while texts and not texts[-1]:
                texts.pop()
            score = (nonempty, len(texts))
            if best is None or score > (best[0], len(best[1])):
                best = (nonempty, texts)
    return best[1] if best else None


def sheet_name_matches_target(sheet_name: str, target_key: str) -> bool:
    n = _norm(sheet_name)
    kn = _norm(target_key)
    return bool(kn and kn in n)


def load_sheet_header_slugs_norm_to_slug(yaml_path: Path | None = None) -> dict[str, str]:
    """归一化后的表头文本 -> sheet_header_slugs 中的英文 slug。"""
    from config.field_mapping import get_field_mapping_yaml_path

    path = yaml_path or get_field_mapping_yaml_path()
    if not path.exists():
        return {}
    try:
        raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    except Exception:
        return {}
    if not isinstance(raw, dict):
        return {}
    block = raw.get("sheet_header_slugs")
    if not isinstance(block, dict):
        return {}
    out: dict[str, str] = {}
    for k, v in block.items():
        if not isinstance(k, str) or not isinstance(v, str):
            continue
        ks, vs = k.strip(), v.strip()
        if not ks or not vs:
            continue
        out[_norm(ks)] = vs
    return out


def analyze_sheet_headers(
    sheet_name: str,
    columns: list[Any],
    *,
    slug_norm_to_slug: dict[str, str] | None = None,
) -> dict[str, Any]:
    """
    返回：
    - unmapped_headers: 非空表头中，既未被字段推断占用、也不在 sheet_header_slugs、也不是任一已知别名精确归一值 的列
    - suggested_slugs: 表头 -> 建议英文 key（仅针对 unmapped）
    """
    slug_map = slug_norm_to_slug if slug_norm_to_slug is not None else load_sheet_header_slugs_norm_to_slug()

    picked = _infer_field_columns(columns, sheet_name)
    assigned_original = {str(v).strip() for v in picked.values() if v}
    assigned_norm = {_normalize_colname(x) for x in assigned_original}

    fm = get_field_mapping_for_sheet(sheet_name)
    alias_norms: set[str] = set()
    for aliases in fm.values():
        for a in aliases:
            alias_norms.add(_normalize_colname(a))

    taken_slugs: set[str] = set(fm.keys()) | set(slug_map.values())

    unmapped: list[str] = []
    suggested: dict[str, str] = {}
    seen: set[str] = set()

    for c in columns:
        raw = str(c).strip() if c is not None else ""
        if not raw:
            continue
        nc = _normalize_colname(raw)
        if nc in seen:
            continue
        seen.add(nc)

        if raw in assigned_original or nc in assigned_norm:
            continue
        if nc in slug_map:
            continue
        if nc in alias_norms:
            continue

        unmapped.append(raw)
        sug = suggest_field_slug(raw, taken_slugs=taken_slugs)
        suggested[raw] = sug
        taken_slugs.add(sug)

    return {
        "sheet_name": sheet_name,
        "unmapped_headers": unmapped,
        "suggested_slugs": suggested,
    }


def analyze_excel_path(
    path: Path,
    target_keys: list[str],
    *,
    slug_norm_to_slug: dict[str, str] | None = None,
) -> dict[str, Any]:
    """扫描工作簿中「表名命中任一 target_key」的 sheet，返回每表分析结果。"""
    keys = [str(k).strip() for k in target_keys if str(k).strip()]
    if not keys:
        return {"ok": False, "error": {"message": "target_keys 为空"}}

    slug_map = slug_norm_to_slug if slug_norm_to_slug is not None else load_sheet_header_slugs_norm_to_slug()

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return {
            "ok": False,
            "error": {
                "message": "无法打开 Excel（文件损坏或格式异常）",
                "exception_type": type(exc).__name__,
                "detail": str(exc),
            },
        }

    sheets_out: list[dict[str, Any]] = []
    try:
        for sn in wb.sheetnames:
            matched = [tk for tk in keys if sheet_name_matches_target(sn, tk)]
            if not matched:
                continue
            ws = wb[sn]
            hdr = _detect_header_row(ws)
            if not hdr:
                sheets_out.append(
                    {
                        "sheet_name": sn,
                        "matched_target_keys": matched,
                        "header_row_found": False,
                        "unmapped_headers": [],
                        "suggested_slugs": {},
                    }
                )
                continue
            part = analyze_sheet_headers(sn, hdr, slug_norm_to_slug=slug_map)
            part["matched_target_keys"] = matched
            part["header_row_found"] = True
            sheets_out.append(part)
    finally:
        wb.close()

    return {
        "ok": True,
        "file": str(path),
        "sheets": sheets_out,
        "update_hint": (
            "可将「表头文本 -> 建议 slug」追加到 config/field_mapping.yaml 的 sheet_header_slugs；"
            "或对有样例的目录运行：python tools/sync_sheet_block_from_samples.py --data <样例目录> --all-sheet-keys"
        ),
    }
