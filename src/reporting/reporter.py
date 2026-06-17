"""Word 审计报告生成（python-docx）。"""

from __future__ import annotations

import logging
from datetime import date, datetime
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

DEFAULT_CHAPTERS: dict[str, bool] = {
    "overview": True,
    "structure": True,
    "supplier": True,
    "audit_flags": True,
    "flags_track": True,
    "related": True,
    "compare": True,
    "supplier_new": False,
    "trade_relationships": False,
    "tax_in_out_deviation": False,
    "finance_reconcile": False,
    "data_quality_summary": False,
    "tax_code_analysis": False,
    "tax_risk_exposure": False,
    "goods_category": False,
    "red_offset_analysis": False,
    "counterparty_risk": False,
    "invoice_timing": False,
    "year_over_year": False,
}


def _chapter_enabled(chapters: dict[str, bool] | None, key: str) -> bool:
    if chapters is None:
        return DEFAULT_CHAPTERS.get(key, True)
    return bool(chapters.get(key, DEFAULT_CHAPTERS.get(key, True)))


def _fmt_amount(v: float | None) -> str:
    if v is None:
        return "—"
    return f"{v:,.2f}"


def _fmt_pct(v: float | None) -> str:
    if v is None:
        return "—"
    return f"{v * 100:.2f}%"


def _add_heading(doc: Any, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _add_kv_table(doc: Any, rows: list[tuple[str, str]]) -> None:
    if not rows:
        doc.add_paragraph("（无数据）")
        return
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v


def _query_overview(conn: Any, stat_year: int) -> dict[str, Any]:
    row = conn.execute(
        """
        SELECT
            coalesce(sum(net_jshj), 0),
            coalesce(sum(normal_cnt + red_cnt + cancel_cnt), 0)::BIGINT,
            coalesce(sum(cancel_cnt), 0)::BIGINT
        FROM dws_inv_trend WHERE stat_year = ?
        """,
        [stat_year],
    ).fetchone()
    sup = conn.execute(
        "SELECT count(DISTINCT supplier_id)::BIGINT FROM dws_sup_conc WHERE stat_year = ?",
        [stat_year],
    ).fetchone()
    from src.audit.config_loader import group_id_for_year

    gid = group_id_for_year(stat_year)
    flags = conn.execute(
        "SELECT count(*)::BIGINT FROM dm_audit_flag WHERE group_id = ?",
        [gid],
    ).fetchone()
    entity_cnt = conn.execute(
        "SELECT count(DISTINCT entity_id)::BIGINT FROM dws_inv_trend WHERE stat_year = ?",
        [stat_year],
    ).fetchone()
    return {
        "total_net": float(row[0] or 0) if row else 0.0,
        "invoice_cnt": int(row[1] or 0) if row else 0,
        "cancel_cnt": int(row[2] or 0) if row else 0,
        "supplier_cnt": int(sup[0] or 0) if sup else 0,
        "flag_cnt": int(flags[0] or 0) if flags else 0,
        "entity_cnt": int(entity_cnt[0] or 0) if entity_cnt else 0,
    }


def _query_trend_rows(conn: Any, stat_year: int) -> list[tuple[int, float]]:
    rows = conn.execute(
        """
        SELECT stat_month, coalesce(sum(net_jshj), 0)
        FROM dws_inv_trend
        WHERE stat_year = ?
        GROUP BY stat_month
        ORDER BY stat_month
        """,
        [stat_year],
    ).fetchall()
    return [(int(r[0]), float(r[1] or 0)) for r in rows or []]


def _query_top_suppliers(conn: Any, stat_year: int, limit: int = 10) -> list[tuple[str, str, float, float]]:
    rows = conn.execute(
        """
        SELECT entity_name, supplier_name, net_jshj, amount_ratio
        FROM dws_sup_conc
        WHERE stat_year = ?
        ORDER BY net_jshj DESC NULLS LAST
        LIMIT ?
        """,
        [stat_year, limit],
    ).fetchall()
    return [
        (str(r[0] or ""), str(r[1] or ""), float(r[2] or 0), float(r[3] or 0))
        for r in rows or []
    ]


def _query_confirmed_flags(conn: Any, stat_year: int, limit: int = 50) -> list[tuple[str, str, str, str, str, str]]:
    from src.audit.config_loader import group_id_for_year

    gid = group_id_for_year(stat_year)
    rows = conn.execute(
        """
        SELECT rule_id, risk_level, entity_name, seller_name, description, coalesce(confirm_note, '')
        FROM dm_audit_flag
        WHERE group_id = ? AND COALESCE(is_confirmed, FALSE) = TRUE
        ORDER BY
            CASE risk_level WHEN '高风险' THEN 1 WHEN '中风险' THEN 2 ELSE 3 END,
            amount DESC NULLS LAST
        LIMIT ?
        """,
        [gid, limit],
    ).fetchall()
    return [
        (str(r[0] or ""), str(r[1] or ""), str(r[2] or ""), str(r[3] or ""), str(r[4] or ""), str(r[5] or ""))
        for r in rows or []
    ]


def _query_flags_by_rule(conn: Any, stat_year: int) -> list[tuple[str, int, int]]:
    from src.audit.config_loader import group_id_for_year

    gid = group_id_for_year(stat_year)
    rows = conn.execute(
        """
        SELECT rule_id,
               count(*)::BIGINT,
               sum(CASE WHEN COALESCE(is_confirmed, FALSE) THEN 1 ELSE 0 END)::BIGINT
        FROM dm_audit_flag
        WHERE group_id = ?
        GROUP BY rule_id
        ORDER BY count(*) DESC, rule_id
        """,
        [gid],
    ).fetchall()
    return [(str(r[0] or ""), int(r[1] or 0), int(r[2] or 0)) for r in rows or []]


def _query_supplier_new(conn: Any, stat_year: int, limit: int = 20) -> list[tuple[str, str, float]]:
    rows = conn.execute(
        """
        SELECT entity_name, supplier_name, net_jshj
        FROM dws_sup_churn
        WHERE stat_year = ? AND churn_kind = 'new'
        ORDER BY net_jshj DESC NULLS LAST
        LIMIT ?
        """,
        [stat_year, limit],
    ).fetchall()
    return [(str(r[0] or ""), str(r[1] or ""), float(r[2] or 0)) for r in rows or []]


def _query_goods_category_top(conn: Any, stat_year: int, limit: int = 20) -> list[tuple[str, str, str, float, int]]:
    try:
        rows = conn.execute(
            """
            SELECT entity_id, tax_code_short, tax_code_level2, sum(net_jshj), sum(invoice_cnt)::BIGINT
            FROM dws_goods_cat
            WHERE stat_year = ?
            GROUP BY entity_id, tax_code_short, tax_code_level2
            ORDER BY sum(net_jshj) DESC NULLS LAST
            LIMIT ?
            """,
            [stat_year, limit],
        ).fetchall()
        return [
            (str(r[0] or ""), str(r[1] or ""), str(r[2] or ""), float(r[3] or 0), int(r[4] or 0))
            for r in rows or []
        ]
    except Exception:
        return []


def _query_red_offset_summary(conn: Any, stat_year: int, limit: int = 15) -> list[tuple[str, int, int, float]]:
    try:
        rows = conn.execute(
            """
            SELECT coalesce(h.gfsbh, h.xfsbh, ''),
                   count(*)::BIGINT,
                   sum(CASE WHEN coalesce(h.fpzt, '') LIKE '%红%' OR coalesce(h.is_orphan_red, FALSE) THEN 1 ELSE 0 END)::BIGINT,
                   coalesce(sum(abs(coalesce(h.net_jshj, h.jshj, 0))) FILTER (WHERE coalesce(h.fpzt, '') LIKE '%红%' OR coalesce(h.is_orphan_red, FALSE)), 0)
            FROM dwd_inv_header h
            WHERE h.stat_year = ?
            GROUP BY 1
            HAVING sum(CASE WHEN coalesce(h.fpzt, '') LIKE '%红%' OR coalesce(h.is_orphan_red, FALSE) THEN 1 ELSE 0 END) > 0
            ORDER BY 3 DESC
            LIMIT ?
            """,
            [stat_year, limit],
        ).fetchall()
        return [
            (str(r[0] or ""), int(r[1] or 0), int(r[2] or 0), float(r[3] or 0))
            for r in rows or []
        ]
    except Exception:
        return []


def _query_counterparty_risk_top(
    conn: Any, stat_year: int, limit: int = 20
) -> list[tuple[str, str, str, float, int, float]]:
    try:
        from src.local_api.counterparty_risk_api import api_dws_counterparty_risk_list

        entity_rows = conn.execute(
            """
            SELECT DISTINCT entity_id
            FROM dm_audit_flag
            WHERE group_id = ?
            UNION
            SELECT DISTINCT entity_id FROM dws_trade_sum WHERE stat_year = ?
            """,
            [f"Y{stat_year}", stat_year],
        ).fetchall()
        scored: list[tuple[str, str, str, float, int, float]] = []
        for (eid_raw,) in entity_rows or []:
            eid = str(eid_raw or "").strip()
            if not eid:
                continue
            res = api_dws_counterparty_risk_list(conn, stat_year=str(stat_year), entity_id=eid, limit=3)
            if not res.get("ok"):
                continue
            for row in res.get("rows") or []:
                scored.append(
                    (
                        eid,
                        str(row.get("counterparty_id") or ""),
                        str(row.get("counterparty_name") or row.get("counterparty_id") or ""),
                        float(row.get("risk_score") or 0),
                        int(row.get("flag_count") or 0),
                        float(row.get("trade_amount") or 0),
                    )
                )
        scored.sort(key=lambda x: (-x[3], -x[5], x[1]))
        return scored[:limit]
    except Exception:
        return []


def _query_trade_top(conn: Any, stat_year: int, limit: int = 20) -> list[tuple[str, str, str, float]]:
    rows = conn.execute(
        """
        SELECT entity_name, counterparty_name, counterparty_role, total_amount
        FROM dws_trade_sum
        WHERE stat_year = ?
        ORDER BY abs(total_amount) DESC NULLS LAST
        LIMIT ?
        """,
        [stat_year, limit],
    ).fetchall()
    return [
        (str(r[0] or ""), str(r[1] or ""), str(r[2] or ""), float(r[3] or 0))
        for r in rows or []
    ]


def _query_tax_deviation_summary(conn: Any, stat_year: int, limit: int = 10) -> list[tuple[str, str, float, float]]:
    """各主体进销结构偏离度（需 DWS 已刷新）。"""
    try:
        rows = conn.execute(
            """
            SELECT entity_id, any_value(entity_name), 0.0, 0.0
            FROM dws_inv_trend
            WHERE stat_year = ?
            GROUP BY entity_id
            LIMIT ?
            """,
            [stat_year, limit],
        ).fetchall()
        out: list[tuple[str, str, float, float]] = []
        for r in rows or []:
            eid = str(r[0] or "")
            ename = str(r[1] or eid)
            dev = None
            try:
                from src.local_api.dws_dashboard_api import api_dws_tax_in_out_deviation

                dev = api_dws_tax_in_out_deviation(conn, stat_year=str(stat_year), entity_id=eid)
            except Exception:
                pass
            if dev and dev.get("ok"):
                out.append(
                    (
                        ename,
                        eid,
                        float(dev.get("mix_deviation_l1") or 0),
                        float(dev.get("exceeded_bucket_count") or 0),
                    )
                )
        out.sort(key=lambda x: x[2], reverse=True)
        return out[:limit]
    except Exception:
        return []


def _query_finance_reconcile_summary(
    conn: Any, stat_year: int, limit: int = 15
) -> tuple[dict[str, dict[str, Any]], list[tuple[str, str, str, str, float, float, float]]]:
    """财务账票核对差异汇总（按最新导入批次）。"""
    try:
        from src.local_api.finance_reconcile_api import _build_reconcile_rows

        batch_row = conn.execute(
            """
            SELECT batch_id
            FROM dm_finance_ledger_batch
            WHERE stat_year = ?
            ORDER BY imported_at DESC NULLS LAST
            LIMIT 1
            """,
            [stat_year],
        ).fetchone()
        if not batch_row or not batch_row[0]:
            return {}, []
        batch_id = str(batch_row[0])
        rows = _build_reconcile_rows(conn, batch_id=batch_id, stat_year=stat_year)
        summary: dict[str, dict[str, Any]] = {
            c: {"count": 0, "diff_amount": 0.0} for c in ("A", "B", "C", "D")
        }
        top: list[tuple[str, str, str, str, float, float, float]] = []
        for r in rows:
            dt = str(r.get("diff_type") or "")
            if dt == "MATCH":
                continue
            if dt in summary:
                summary[dt]["count"] += 1
                summary[dt]["diff_amount"] = round(
                    summary[dt]["diff_amount"] + abs(float(r.get("diff_amount") or 0)), 2
                )
            top.append(
                (
                    str(r.get("entity_name") or r.get("tax_id") or ""),
                    str(r.get("tax_id") or ""),
                    dt,
                    str(r.get("subject_name") or r.get("subject_code") or "—"),
                    float(r.get("ledger_amount") or 0) if r.get("ledger_amount") is not None else 0.0,
                    float(r.get("invoice_net") or 0) if r.get("invoice_net") is not None else 0.0,
                    abs(float(r.get("diff_amount") or 0)),
                )
            )
        top.sort(key=lambda x: x[6], reverse=True)
        return summary, top[:limit]
    except Exception:
        logger.exception("finance reconcile report summary")
        return {}, []


def _query_flags(conn: Any, stat_year: int, limit: int = 50) -> list[tuple[str, str, str, str, str]]:
    from src.audit.config_loader import group_id_for_year

    gid = group_id_for_year(stat_year)
    rows = conn.execute(
        """
        SELECT rule_id, risk_level, entity_name, seller_name, description
        FROM dm_audit_flag
        WHERE group_id = ?
        ORDER BY
            CASE risk_level WHEN '高风险' THEN 1 WHEN '中风险' THEN 2 ELSE 3 END,
            amount DESC NULLS LAST
        LIMIT ?
        """,
        [gid, limit],
    ).fetchall()
    return [
        (str(r[0] or ""), str(r[1] or ""), str(r[2] or ""), str(r[3] or ""), str(r[4] or ""))
        for r in rows or []
    ]


def _query_data_quality_summary(conn: Any, stat_year: int) -> tuple[dict[str, Any], dict[str, Any]] | None:
    """导入质量域 KPI 与分域指标（轻量摘要，完整明细见交付包 CSV）。"""
    try:
        from src.local_api.data_quality import load_dq_domain_overview

        dq = load_dq_domain_overview(conn, stat_year=stat_year)
        if not dq.get("ok"):
            return None
        return dict(dq.get("kpi") or {}), dict(dq.get("domains") or {})
    except Exception:
        logger.exception("data_quality report summary")
        return None


def _query_tax_code_summary(conn: Any, stat_year: int) -> dict[str, Any] | None:
    """税收分类编码覆盖率摘要。"""
    try:
        from src.local_api.tax_code_analysis_api import api_tax_code_analysis_overview

        payload = api_tax_code_analysis_overview(conn, stat_year=str(stat_year))
        if not payload.get("ok"):
            return None
        return dict(payload.get("data") or payload)
    except Exception:
        logger.exception("tax_code report summary")
        return None


def _query_tax_risk_summary(conn: Any, stat_year: int, limit: int = 15) -> list[tuple[str, str, float, float, float]]:
    """税风险敞口 Top 主体摘要。"""
    try:
        from src.local_api.dws_dashboard_api import api_dws_tax_risk_exposure

        rows = conn.execute(
            """
            SELECT trim(COALESCE(entity_id, '')) AS eid,
                   max(trim(COALESCE(entity_name, ''))) AS ename
            FROM ads_scorecard
            WHERE stat_year = ? AND length(trim(COALESCE(entity_id, ''))) > 0
            GROUP BY 1
            ORDER BY 1
            LIMIT ?
            """,
            [stat_year, limit],
        ).fetchall()
        out: list[tuple[str, str, float, float, float]] = []
        for r in rows or []:
            eid, ename = str(r[0] or ""), str(r[1] or "")
            if not eid:
                continue
            payload = api_dws_tax_risk_exposure(conn, stat_year=str(stat_year), entity_id=eid)
            if not payload.get("ok"):
                continue
            out.append(
                (
                    ename or eid,
                    eid,
                    float(payload.get("total_exposure") or 0),
                    float(payload.get("deviation_exposure") or 0),
                    float(payload.get("high_risk_coding_amount") or 0),
                )
            )
        out.sort(key=lambda x: x[2], reverse=True)
        return out[:limit]
    except Exception:
        logger.exception("tax_risk report summary")
        return []


def _query_related(conn: Any, stat_year: int) -> tuple[list[tuple], list[tuple]]:
    from src.audit.config_loader import group_id_for_year

    gid = group_id_for_year(stat_year)
    circ = conn.execute(
        """
        SELECT party_a_name, party_b_name, amount_a_to_b, amount_b_to_a, circular_ratio, risk_level
        FROM dm_circ_inv WHERE group_id = ?
        ORDER BY circular_ratio DESC NULLS LAST
        LIMIT 20
        """,
        [gid],
    ).fetchall()
    shell = conn.execute(
        """
        SELECT group_member_name, intermediary_name, amount_in, amount_out, passthrough_ratio, risk_level
        FROM dm_shell_co WHERE group_id = ?
        ORDER BY passthrough_ratio DESC NULLS LAST
        LIMIT 20
        """,
        [gid],
    ).fetchall()
    return list(circ or []), list(shell or [])


def _query_scorecard(conn: Any, stat_year: int) -> list[tuple]:
    from src.audit.config_loader import group_id_for_year

    gid = group_id_for_year(stat_year)
    rows = conn.execute(
        """
        SELECT entity_name, total_amount, total_count, supplier_count,
               flag_total, flag_high, risk_score, risk_level
        FROM ads_scorecard
        WHERE group_id = ? AND stat_year = ?
        ORDER BY risk_score ASC, flag_high DESC
        LIMIT 50
        """,
        [gid, stat_year],
    ).fetchall()
    return list(rows or [])


def generate_audit_report_docx(
    conn: Any,
    *,
    stat_year: int,
    title: str,
    chapters: dict[str, bool] | None = None,
    output_dir: str,
) -> dict[str, Any]:
    """生成 Word 审计报告，返回文件路径信息。"""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        return {
            "ok": False,
            "error": {
                "message": "缺少 python-docx 依赖，请安装 requirements.txt",
                "exception_type": type(exc).__name__,
            },
        }

    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_title = "".join(c if c.isalnum() or c in "._-" else "_" for c in title[:40])
    file_name = f"audit_report_{stat_year}_{ts}_{safe_title}.docx"
    file_path = out / file_name

    doc = Document()
    overview = _query_overview(conn, stat_year)

    # 封面
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = run.font.size  # keep default
    doc.add_paragraph("")
    cover = doc.add_paragraph(f"分析年度：{stat_year}")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(
        f"数据来源：{overview['entity_cnt']} 家主体，共 {overview['invoice_cnt']:,} 张发票"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gen = doc.add_paragraph(f"生成日期：{date.today().isoformat()}")
    gen.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("保密等级：内部资料").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    if _chapter_enabled(chapters, "overview"):
        _add_heading(doc, "第一章  数据概览", 1)
        _add_kv_table(
            doc,
            [
                ("分析年度", str(stat_year)),
                ("净额合计（net_jshj）", _fmt_amount(overview["total_net"])),
                ("发票总张数", f"{overview['invoice_cnt']:,}"),
                ("涉及供应商数", str(overview["supplier_cnt"])),
                ("作废/冲红相关张数", str(overview["cancel_cnt"])),
                ("审计疑点数", str(overview["flag_cnt"])),
                ("涉及主体数", str(overview["entity_cnt"])),
            ],
        )
        doc.add_paragraph(
            "说明：金额口径来自 DWS 聚合层（dws_inv_trend）；疑点来自 dm_audit_flag。"
        )

    if _chapter_enabled(chapters, "structure"):
        _add_heading(doc, "第二章  发票结构分析", 1)
        trend_rows = _query_trend_rows(conn, stat_year)
        if trend_rows:
            table = doc.add_table(rows=len(trend_rows) + 1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "月份"
            table.rows[0].cells[1].text = "净额合计"
            for i, (m, amt) in enumerate(trend_rows, start=1):
                table.rows[i].cells[0].text = f"{stat_year}-{m:02d}"
                table.rows[i].cells[1].text = _fmt_amount(amt)
        else:
            doc.add_paragraph("（暂无月度趋势数据，请先刷新 DWS 聚合。）")

    if _chapter_enabled(chapters, "supplier"):
        _add_heading(doc, "第三章  供应商分析", 1)
        tops = _query_top_suppliers(conn, stat_year)
        if tops:
            table = doc.add_table(rows=len(tops) + 1, cols=4)
            table.style = "Table Grid"
            hdr = ("购方企业", "供应商", "采购净额", "金额占比")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, (ent, sup, amt, ratio) in enumerate(tops, start=1):
                table.rows[i].cells[0].text = ent or "—"
                table.rows[i].cells[1].text = sup or "—"
                table.rows[i].cells[2].text = _fmt_amount(amt)
                table.rows[i].cells[3].text = _fmt_pct(ratio)
        else:
            doc.add_paragraph("（暂无供应商集中度数据。）")

    if _chapter_enabled(chapters, "audit_flags"):
        _add_heading(doc, "第四章  审计疑点清单", 1)
        rule_dist = _query_flags_by_rule(conn, stat_year)
        if rule_dist:
            doc.add_heading("4.1 按规则分布", level=2)
            table = doc.add_table(rows=len(rule_dist) + 1, cols=3)
            table.style = "Table Grid"
            hdr = ("规则", "疑点总数", "已确认")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, (rule, total_n, confirmed_n) in enumerate(rule_dist, start=1):
                table.rows[i].cells[0].text = rule
                table.rows[i].cells[1].text = str(total_n)
                table.rows[i].cells[2].text = str(confirmed_n)
        flags = _query_flags(conn, stat_year)
        doc.add_heading("4.2 疑点明细（Top）", level=2)
        if flags:
            table = doc.add_table(rows=len(flags) + 1, cols=5)
            table.style = "Table Grid"
            hdr = ("规则", "风险", "企业", "供应商", "描述")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, (rule, risk, ent, seller, desc) in enumerate(flags, start=1):
                table.rows[i].cells[0].text = rule
                table.rows[i].cells[1].text = risk
                table.rows[i].cells[2].text = ent or "—"
                table.rows[i].cells[3].text = seller or "—"
                table.rows[i].cells[4].text = (desc or "")[:200]
        else:
            doc.add_paragraph("（当前年度暂无审计疑点，请先运行疑点扫描。）")

    if _chapter_enabled(chapters, "flags_track"):
        _add_heading(doc, "附录  已确认疑点摘要", 1)
        confirmed = _query_confirmed_flags(conn, stat_year)
        if confirmed:
            table = doc.add_table(rows=len(confirmed) + 1, cols=6)
            table.style = "Table Grid"
            hdr = ("规则", "风险", "企业", "供应商", "描述", "跟踪说明")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, (rule, risk, ent, seller, desc, note) in enumerate(confirmed, start=1):
                table.rows[i].cells[0].text = rule
                table.rows[i].cells[1].text = risk
                table.rows[i].cells[2].text = ent or "—"
                table.rows[i].cells[3].text = seller or "—"
                table.rows[i].cells[4].text = (desc or "")[:160]
                table.rows[i].cells[5].text = (note or "")[:160]
        else:
            doc.add_paragraph("（当前年度暂无已确认疑点。）")

    if _chapter_enabled(chapters, "supplier_new"):
        _add_heading(doc, "专题  新增供应商", 1)
        news = _query_supplier_new(conn, stat_year)
        if news:
            table = doc.add_table(rows=len(news) + 1, cols=3)
            table.style = "Table Grid"
            hdr = ("购方企业", "新增供应商", "采购净额")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, (ent, sup, amt) in enumerate(news, start=1):
                table.rows[i].cells[0].text = ent or "—"
                table.rows[i].cells[1].text = sup or "—"
                table.rows[i].cells[2].text = _fmt_amount(amt)
        else:
            doc.add_paragraph("（暂无新增供应商数据。）")

    if _chapter_enabled(chapters, "trade_relationships"):
        _add_heading(doc, "专题  往来关系摘要", 1)
        trades = _query_trade_top(conn, stat_year)
        if trades:
            table = doc.add_table(rows=len(trades) + 1, cols=4)
            table.style = "Table Grid"
            hdr = ("主体", "对方", "角色", "金额")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, (ent, cp, role, amt) in enumerate(trades, start=1):
                table.rows[i].cells[0].text = ent or "—"
                table.rows[i].cells[1].text = cp or "—"
                table.rows[i].cells[2].text = role or "—"
                table.rows[i].cells[3].text = _fmt_amount(amt)
        else:
            doc.add_paragraph("（暂无往来关系数据。）")

    if _chapter_enabled(chapters, "tax_in_out_deviation"):
        _add_heading(doc, "专题  进销偏离分析", 1)
        dev_rows = _query_tax_deviation_summary(conn, stat_year)
        if dev_rows:
            table = doc.add_table(rows=len(dev_rows) + 1, cols=4)
            table.style = "Table Grid"
            hdr = ("企业", "税号", "结构偏离度", "超阈值档位数")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, (ename, eid, mix, exc) in enumerate(dev_rows, start=1):
                table.rows[i].cells[0].text = ename
                table.rows[i].cells[1].text = eid
                table.rows[i].cells[2].text = f"{mix:.4f}"
                table.rows[i].cells[3].text = str(int(exc))
        else:
            doc.add_paragraph("（暂无进销偏离分析数据。）")

    if _chapter_enabled(chapters, "finance_reconcile"):
        _add_heading(doc, "专题  财务账票核对差异", 1)
        fin_summary, fin_top = _query_finance_reconcile_summary(conn, stat_year)
        if fin_summary and any(fin_summary[c]["count"] > 0 for c in fin_summary):
            doc.add_paragraph(
                "以下汇总基于该年度最近导入的财务账表批次与 DWD 发票净额比对结果；"
                "差异类型 A=票面有账面无，B=账面有票面无，C=双方有金额不一致，D=待分类。"
            )
            type_rows = [
                ("A 票面有账面无", str(fin_summary["A"]["count"]), _fmt_amount(fin_summary["A"]["diff_amount"])),
                ("B 账面有票面无", str(fin_summary["B"]["count"]), _fmt_amount(fin_summary["B"]["diff_amount"])),
                ("C 金额不一致", str(fin_summary["C"]["count"]), _fmt_amount(fin_summary["C"]["diff_amount"])),
                ("D 待分类", str(fin_summary["D"]["count"]), _fmt_amount(fin_summary["D"]["diff_amount"])),
            ]
            table = doc.add_table(rows=len(type_rows) + 1, cols=3)
            table.style = "Table Grid"
            for j, h in enumerate(("差异类型", "笔数", "差异金额合计")):
                table.rows[0].cells[j].text = h
            for i, row in enumerate(type_rows, start=1):
                for j, val in enumerate(row):
                    table.rows[i].cells[j].text = val
            if fin_top:
                doc.add_heading("差异金额 Top 明细", level=2)
                table2 = doc.add_table(rows=len(fin_top) + 1, cols=7)
                table2.style = "Table Grid"
                hdr = ("主体", "税号", "类型", "科目", "账表金额", "发票净额", "差异金额")
                for j, h in enumerate(hdr):
                    table2.rows[0].cells[j].text = h
                for i, (ename, eid, dt, subj, led, inv, diff) in enumerate(fin_top, start=1):
                    table2.rows[i].cells[0].text = ename or "—"
                    table2.rows[i].cells[1].text = eid or "—"
                    table2.rows[i].cells[2].text = dt
                    table2.rows[i].cells[3].text = subj or "—"
                    table2.rows[i].cells[4].text = _fmt_amount(led) if led else "—"
                    table2.rows[i].cells[5].text = _fmt_amount(inv) if inv else "—"
                    table2.rows[i].cells[6].text = _fmt_amount(diff)
        else:
            doc.add_paragraph("（暂无财务账票核对差异数据，请先导入账表并完成核对。）")

    if _chapter_enabled(chapters, "data_quality_summary"):
        _add_heading(doc, "专题  数据质量域摘要", 1)
        dq_pair = _query_data_quality_summary(conn, stat_year)
        if dq_pair:
            kpi, domains = dq_pair
            _add_kv_table(
                doc,
                [
                    ("已扫描票头", str(kpi.get("scanned_headers") or 0)),
                    ("异常票头", str(kpi.get("anomaly_headers") or 0)),
                    ("阻塞级", str(kpi.get("block_count") or 0)),
                    ("警告级", str(kpi.get("warn_count") or 0)),
                    ("提示级", str(kpi.get("info_count") or 0)),
                ],
            )
            domain_rows: list[tuple[str, str, str]] = []
            label_map = {
                "uniqueness": "身份与键·唯一性",
                "tax_id": "购销方识别号",
                "cross_table": "跨表对齐",
                "header_detail": "头明细一致",
                "semantic": "语义与专项",
                "red_link": "红冲关联",
            }
            for key, label in label_map.items():
                block = domains.get(key) or {}
                if not isinstance(block, dict):
                    continue
                primary = next((str(v) for k, v in block.items() if k.endswith("_count") or k.endswith("_tickets")), "—")
                if primary != "—" and primary not in ("0", "None"):
                    domain_rows.append((label, primary, "见交付包 data_quality CSV"))
            if domain_rows:
                doc.add_heading("分域指标（节选）", level=2)
                table = doc.add_table(rows=len(domain_rows) + 1, cols=3)
                table.style = "Table Grid"
                for j, h in enumerate(("质量域", "指标值", "说明")):
                    table.rows[0].cells[j].text = h
                for i, row in enumerate(domain_rows, start=1):
                    for j, val in enumerate(row):
                        table.rows[i].cells[j].text = val
            doc.add_paragraph(
                "说明：完整分域指标已写入交付包 exports/data_quality_domain_summary_{year}.csv。".format(
                    year=stat_year
                )
            )
        else:
            doc.add_paragraph("（暂无数据质量扫描结果，请先完成导入质量评估。）")

    if _chapter_enabled(chapters, "tax_code_analysis"):
        _add_heading(doc, "专题  税收分类编码分析", 1)
        tax = _query_tax_code_summary(conn, stat_year)
        if tax:
            _add_kv_table(
                doc,
                [
                    ("含编码明细行", str(tax.get("lines_with_code") or 0)),
                    ("维表命中行", str(tax.get("matched_lines") or 0)),
                    ("未匹配行", str(tax.get("unmatched_lines") or 0)),
                    ("命中率", _fmt_pct(float(tax.get("match_rate") or 0))),
                    ("未匹配金额", _fmt_amount(float(tax.get("unmatched_amount") or 0))),
                    ("HIGH 类目税额占比", _fmt_pct(float(tax.get("high_risk_amount_share") or 0))),
                ],
            )
            if tax.get("top_category_name"):
                doc.add_paragraph(
                    f"Top1 二级类目：{tax.get('top_category_name')}（占比 {_fmt_pct(float(tax.get('top_category_share') or 0))}）"
                )
        else:
            doc.add_paragraph("（暂无税收分类编码分析数据，请先维护 dim_tax_code 并完成 DWD 落盘。）")

    if _chapter_enabled(chapters, "tax_risk_exposure"):
        _add_heading(doc, "专题  税风险敞口", 1)
        risk_rows = _query_tax_risk_summary(conn, stat_year)
        if risk_rows:
            doc.add_paragraph(
                "敞口口径：超阈值进销偏离档位金额差 + 高风险税收分类编码税额 + RULE-05/08 疑点金额 + RULE-TAX-DEV；"
                "完整分项见交付包 exports/tax_risk_exposure CSV。"
            )
            table = doc.add_table(rows=len(risk_rows) + 1, cols=5)
            table.style = "Table Grid"
            hdr = ("企业", "税号", "总敞口", "偏离敞口", "高风险编码税额")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, (ename, eid, total, dev, coding) in enumerate(risk_rows, start=1):
                table.rows[i].cells[0].text = ename
                table.rows[i].cells[1].text = eid
                table.rows[i].cells[2].text = _fmt_amount(total)
                table.rows[i].cells[3].text = _fmt_amount(dev)
                table.rows[i].cells[4].text = _fmt_amount(coding)
        else:
            doc.add_paragraph("（暂无税风险敞口数据，请先完成 DWS 聚合与进销偏离分析。）")

    if _chapter_enabled(chapters, "goods_category"):
        _add_heading(doc, "专题  品类结构分析", 1)
        cat_rows = _query_goods_category_top(conn, stat_year)
        if cat_rows:
            table = doc.add_table(rows=len(cat_rows) + 1, cols=5)
            table.style = "Table Grid"
            hdr = ("主体税号", "税码前缀", "二级前缀", "净额", "行数")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, (eid, short, lvl2, amt, cnt) in enumerate(cat_rows, start=1):
                table.rows[i].cells[0].text = eid or "—"
                table.rows[i].cells[1].text = short or "—"
                table.rows[i].cells[2].text = lvl2 or "—"
                table.rows[i].cells[3].text = _fmt_amount(amt)
                table.rows[i].cells[4].text = str(cnt)
        else:
            doc.add_paragraph("（暂无品类结构数据，请先执行 DWS 品类汇总刷新。）")

    if _chapter_enabled(chapters, "red_offset_analysis"):
        _add_heading(doc, "专题  红冲/作废分析", 1)
        red_rows = _query_red_offset_summary(conn, stat_year)
        if red_rows:
            table = doc.add_table(rows=len(red_rows) + 1, cols=4)
            table.style = "Table Grid"
            hdr = ("主体税号", "发票张数", "红票张数", "红冲金额")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, (eid, header_cnt, red_cnt, red_amt) in enumerate(red_rows, start=1):
                table.rows[i].cells[0].text = eid or "—"
                table.rows[i].cells[1].text = str(header_cnt)
                table.rows[i].cells[2].text = str(red_cnt)
                table.rows[i].cells[3].text = _fmt_amount(red_amt)
        else:
            doc.add_paragraph("（暂无红冲/作废专题数据。）")

    if _chapter_enabled(chapters, "counterparty_risk"):
        _add_heading(doc, "专题  对手风险聚合", 1)
        risk_rows = _query_counterparty_risk_top(conn, stat_year)
        if risk_rows:
            table = doc.add_table(rows=len(risk_rows) + 1, cols=6)
            table.style = "Table Grid"
            hdr = ("主体税号", "对手税号", "对手名称", "风险分", "疑点数", "交易金额")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, (eid, cp_id, cp_name, score, flags, trade) in enumerate(risk_rows, start=1):
                table.rows[i].cells[0].text = eid or "—"
                table.rows[i].cells[1].text = cp_id or "—"
                table.rows[i].cells[2].text = cp_name or "—"
                table.rows[i].cells[3].text = f"{score:.2f}"
                table.rows[i].cells[4].text = str(flags)
                table.rows[i].cells[5].text = _fmt_amount(trade)
        else:
            doc.add_paragraph("（暂无对手风险聚合数据。）")

    if _chapter_enabled(chapters, "invoice_timing"):
        _add_heading(doc, "专题  开票时间行为", 1)
        doc.add_paragraph("（开票时间行为专题请在线查看「开票时间行为」分析页；报告导出暂未嵌入明细表。）")

    if _chapter_enabled(chapters, "year_over_year"):
        _add_heading(doc, "专题  跨年结构对比", 1)
        doc.add_paragraph("（跨年结构对比专题请在线查看「跨年对比」分析页；报告导出暂未嵌入明细表。）")

    if _chapter_enabled(chapters, "related"):
        _add_heading(doc, "第五章  关联交易分析", 1)
        circ, shell = _query_related(conn, stat_year)
        doc.add_heading("5.1 对开发票", level=2)
        if circ:
            table = doc.add_table(rows=len(circ) + 1, cols=6)
            table.style = "Table Grid"
            hdr = ("甲方", "乙方", "A→B金额", "B→A金额", "对开比例", "风险")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, r in enumerate(circ, start=1):
                table.rows[i].cells[0].text = str(r[0] or "")
                table.rows[i].cells[1].text = str(r[1] or "")
                table.rows[i].cells[2].text = _fmt_amount(float(r[2] or 0))
                table.rows[i].cells[3].text = _fmt_amount(float(r[3] or 0))
                table.rows[i].cells[4].text = _fmt_pct(float(r[4] or 0))
                table.rows[i].cells[5].text = str(r[5] or "")
        else:
            doc.add_paragraph("（暂无对开发票记录。）")
        doc.add_heading("5.2 疑似通道公司", level=2)
        if shell:
            table = doc.add_table(rows=len(shell) + 1, cols=6)
            table.style = "Table Grid"
            hdr = ("集团成员", "中间商", "收款", "付款", "穿透比例", "风险")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, r in enumerate(shell, start=1):
                table.rows[i].cells[0].text = str(r[0] or "")
                table.rows[i].cells[1].text = str(r[1] or "")
                table.rows[i].cells[2].text = _fmt_amount(float(r[2] or 0))
                table.rows[i].cells[3].text = _fmt_amount(float(r[3] or 0))
                table.rows[i].cells[4].text = _fmt_pct(float(r[4] or 0))
                table.rows[i].cells[5].text = str(r[5] or "")
        else:
            doc.add_paragraph("（暂无通道公司识别结果。）")

    if _chapter_enabled(chapters, "compare"):
        _add_heading(doc, "第六章  子公司横向对比", 1)
        score_rows = _query_scorecard(conn, stat_year)
        if score_rows:
            table = doc.add_table(rows=len(score_rows) + 1, cols=8)
            table.style = "Table Grid"
            hdr = ("企业", "进项净额", "张数", "供应商数", "疑点", "高风险", "得分", "评级")
            for j, h in enumerate(hdr):
                table.rows[0].cells[j].text = h
            for i, r in enumerate(score_rows, start=1):
                table.rows[i].cells[0].text = str(r[0] or "")
                table.rows[i].cells[1].text = _fmt_amount(float(r[1] or 0))
                table.rows[i].cells[2].text = str(int(r[2] or 0))
                table.rows[i].cells[3].text = str(int(r[3] or 0))
                table.rows[i].cells[4].text = str(int(r[4] or 0))
                table.rows[i].cells[5].text = str(int(r[5] or 0))
                table.rows[i].cells[6].text = f"{float(r[6] or 0):.0f}"
                table.rows[i].cells[7].text = str(r[7] or "")
        else:
            doc.add_paragraph("（暂无评分卡数据，请先刷新 ads_scorecard。）")

    doc.save(str(file_path))
    return {
        "ok": True,
        "file_name": file_name,
        "file_path": str(file_path),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }


def generate_report(output_dir: str) -> str:
    """兼容旧占位接口。"""
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    report = out / "report_placeholder.txt"
    report.write_text(
        "请使用 generate_audit_report_docx 或前端「报告配置」页生成 Word 报告。",
        encoding="utf-8",
    )
    return str(report)
