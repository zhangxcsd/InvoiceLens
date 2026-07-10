"""
组织维度双树（dim_org_node + dim_org_hier）Excel 导入与路径重算。

模板 22 列；第 1 行表头、第 2 行中文说明、第 3 行起为数据。
仍兼容旧版 18/20 列模板（无上级单位名称或国家出资企业列时自动推断）。
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable

import pandas as pd

logger = logging.getLogger(__name__)

AUTOMATION_TASK_CODE = "dim.org_hier.build"
TASK_DISPLAY_NAME = "组织层级双树 · Excel 导入/重算"

ProgressFn = Callable[[str, str], None]

EXPECTED_COLUMNS = [
    "entity_id",
    "entity_fullname",
    "entity_shortname",
    "entity_type",
    "sys_id",
    "stat_year",
    "main_business",
    "industry_id",
    "industry_name",
    "is_stat_inc",
    "state_investor_id",
    "state_investor_name",
    "mg_parent_id",
    "mg_parent_name",
    "mg_sort_no",
    "eq_parent_id",
    "eq_parent_name",
    "eq_sort_no",
    "eq_shareholding_ratio",
    "reg_capital",
    "is_active",
    "hier_diff_note",
]

ENTITY_TYPE_OPTIONS = [
    "有限责任公司",
    "股份有限公司",
    "国有独资公司",
    "全民所有制",
    "集体所有制",
    "合伙企业",
    "个人独资企业",
    "外商投资企业",
    "其他",
]

COLUMN_HINTS_ZH: dict[str, str] = {
    "entity_id": "统一社会信用代码【必填】18位；根节点填 ROOT_{sys_id}",
    "entity_fullname": "企业全称【必填】与工商登记一致",
    "entity_shortname": "企业简称【选填】树节点展示用；无简称时路径自动用全称",
    "entity_type": "企业组织形态【必填】工商登记类型，如有限责任公司、股份有限公司",
    "sys_id": "监管体系代码【必填】系统预置字典项，须从下拉或「监管体系说明」页选择，不可自由填写；维护入口：维度管理→监管体系",
    "stat_year": "统计年度【必填】4位年份；同一文件只允许一个年度",
    "main_business": "主责主业【选填】",
    "industry_id": "行业门类代码【选填】",
    "industry_name": "行业门类名称【选填】",
    "is_stat_inc": "是否纳入统计【必填】填「是」或「否」；根节点填「否」",
    "state_investor_id": "国家出资企业代码【条件必填】填一级企业 entity_id；一级企业一般填自身；根节点留空",
    "state_investor_name": "国家出资企业名称【选填】须与本文件中该代码对应行的 entity_fullname 一致，便于核对",
    "mg_parent_id": "上级管理单位代码【必填】填上级的 entity_id；根节点填自己的 entity_id",
    "mg_parent_name": "上级管理单位名称【选填】须与本文件中该代码对应行的 entity_fullname 一致，便于核对",
    "mg_sort_no": "管理排序号【必填】同一管理上级下从 1 起",
    "eq_parent_id": "上级产权单位代码【必填】填上级的 entity_id；根节点填自己的 entity_id",
    "eq_parent_name": "上级产权单位名称【选填】须与本文件中该代码对应行的 entity_fullname 一致，便于核对",
    "eq_sort_no": "产权排序号【必填】同一产权上级下从 1 起",
    "eq_shareholding_ratio": "直接持股比例【选填】0.0001~1.0000；留空表示全资",
    "reg_capital": "注册资本（万元）【选填】",
    "is_active": "是否在营【必填】填「是」或「否」",
    "hier_diff_note": "管产差异说明【条件必填】管理/产权上级不一致时建议填写",
}

ORG_HIER_TEMPLATE_FILENAME = "组织维度导入模板.xlsx"

_BOOL_YES = {"是", "true", "1", "yes", "y"}
_BOOL_NO = {"否", "false", "0", "no", "n"}


def _project_root() -> Path:
    return Path(__file__).resolve().parents[2]


def _norm_cell(v: Any) -> str | None:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    s = str(v).strip()
    if not s or s.lower() in {"nan", "none", "nat", "<na>"}:
        return None
    return s


def _parse_bool_cn(raw: Any, *, field: str, row_no: int) -> tuple[bool | None, str | None]:
    s = _norm_cell(raw)
    if s is None:
        return None, f"第 {row_no} 行：{field} 不能为空"
    low = s.lower()
    if low in _BOOL_YES or s == "是":
        return True, None
    if low in _BOOL_NO or s == "否":
        return False, None
    return None, f"第 {row_no} 行：{field} 须填「是」或「否」"


def _parse_int(raw: Any, *, field: str, row_no: int, min_val: int = 0) -> tuple[int | None, str | None]:
    s = _norm_cell(raw)
    if s is None:
        return None, f"第 {row_no} 行：{field} 不能为空"
    try:
        val = int(float(s))
    except (TypeError, ValueError):
        return None, f"第 {row_no} 行：{field} 须为整数"
    if val < min_val:
        return None, f"第 {row_no} 行：{field} 须 >= {min_val}"
    return val, None


def _parse_optional_decimal(raw: Any, *, field: str, row_no: int) -> tuple[Decimal | None, str | None]:
    s = _norm_cell(raw)
    if s is None:
        return None, None
    try:
        val = Decimal(s)
    except (InvalidOperation, ValueError):
        return None, f"第 {row_no} 行：{field} 格式无效"
    if val < Decimal("0.0001") or val > Decimal("1.0000"):
        return None, f"第 {row_no} 行：{field} 须在 0.0001~1.0000 之间"
    return val, None


def _parse_optional_date(raw: Any) -> date | None:
    s = _norm_cell(raw)
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def _hier_id(entity_id: str, stat_year: int) -> str:
    return hashlib.md5(f"{entity_id}{stat_year}".encode("utf-8")).hexdigest()


def _valid_entity_id(raw: str) -> bool:
    s = raw.strip()
    if s.startswith("ROOT_"):
        return True
    compact = re.sub(r"[\s-]+", "", s)
    return len(compact) == 18


def _load_active_sys_ids(conn: Any) -> list[str]:
    try:
        rows = conn.execute(
            "SELECT sys_id FROM dim_org_sys WHERE is_active = TRUE ORDER BY sort_no, sys_id"
        ).fetchall()
        sys_ids = [str(r[0]).strip() for r in rows or [] if r[0]]
        return sys_ids
    except Exception as exc:  # noqa: BLE001
        logger.warning("读取 dim_org_sys 失败: %s", exc)
        return []


def _load_sys_catalog(conn: Any) -> list[dict[str, str]]:
    try:
        rows = conn.execute(
            """
            SELECT sys_id, sys_name, admin_level, COALESCE(gov_owner, ''), COALESCE(description, '')
            FROM dim_org_sys
            WHERE is_active = TRUE
            ORDER BY sort_no, sys_id
            """
        ).fetchall()
        return [
            {
                "sys_id": str(r[0] or "").strip(),
                "sys_name": str(r[1] or "").strip(),
                "admin_level": str(r[2] or "").strip(),
                "gov_owner": str(r[3] or "").strip(),
                "description": str(r[4] or "").strip(),
            }
            for r in rows or []
            if r[0]
        ]
    except Exception as exc:  # noqa: BLE001
        logger.warning("读取 dim_org_sys 目录失败: %s", exc)
        return []


def _norm_header_cell(v: Any) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return str(v).strip()


def generate_org_hier_import_template(conn: Any) -> tuple[bytes, str]:
    """生成空白组织维度 Excel 导入模板（表头 + 中文说明行 + 监管体系说明页）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    sys_ids = _load_active_sys_ids(conn)
    sys_catalog = _load_sys_catalog(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "组织维度"

    gray = Font(color="808080")
    for col_idx, col_name in enumerate(EXPECTED_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx, value=col_name)
        hint_cell = ws.cell(row=2, column=col_idx, value=COLUMN_HINTS_ZH.get(col_name, ""))
        hint_cell.font = gray
        width = min(48, max(12, len(col_name) + 2, len(hint_cell.value or "") // 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    bool_dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    ws.add_data_validation(bool_dv)
    bool_dv.add("J3:J10000")
    active_dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    ws.add_data_validation(active_dv)
    active_dv.add("U3:U10000")
    if sys_ids:
        sys_dv = DataValidation(type="list", formula1=f'"{",".join(sys_ids)}"', allow_blank=True)
        ws.add_data_validation(sys_dv)
        sys_dv.add("E3:E10000")
    if ENTITY_TYPE_OPTIONS:
        type_dv = DataValidation(
            type="list",
            formula1=f'"{",".join(ENTITY_TYPE_OPTIONS)}"',
            allow_blank=True,
        )
        ws.add_data_validation(type_dv)
        type_dv.add("D3:D10000")

    ref = wb.create_sheet("监管体系说明")
    ref.append(
        [
            "sys_id（监管体系代码）",
            "sys_name（名称）",
            "admin_level（层级）",
            "gov_owner（主管单位）",
            "description（说明）",
        ]
    )
    for row in sys_catalog:
        ref.append(
            [
                row["sys_id"],
                row["sys_name"],
                row["admin_level"],
                row["gov_owner"],
                row["description"],
            ]
        )
    ref.append([])
    if sys_catalog:
        ref.append(["说明", "sys_id 为系统预置字典项（非自由填写）；导入组织维度时每行须从本页选择对应代码。"])
        ref.append(["维护入口", "维度管理 → 企业组织维度 → 监管体系（可导出最新清单）"])
    else:
        ref.append(["说明", "当前尚无预置监管体系。请先在「维度管理 → 企业组织维度 → 监管体系」维护后再导入。"])
        ref.append(["维护入口", "维度管理 → 企业组织维度 → 监管体系"])
    ref_col_widths = (16, 14, 10, 22, 40)
    for col_idx, width in enumerate(ref_col_widths, start=1):
        ref.column_dimensions[get_column_letter(col_idx)].width = width

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue(), ORG_HIER_TEMPLATE_FILENAME


def _frame_from_org_sheet(raw: pd.DataFrame) -> pd.DataFrame:
    """按第 1 行英文表头映射列；兼容旧版模板（缺 mg/eq_parent_name、state_investor 等列）。"""
    if raw.empty:
        return pd.DataFrame(columns=EXPECTED_COLUMNS)
    headers = [_norm_header_cell(v) for v in raw.iloc[0].tolist()]
    data = raw.iloc[2:].reset_index(drop=True)
    if not any(headers):
        data = data.iloc[:, : len(EXPECTED_COLUMNS)]
        data.columns = EXPECTED_COLUMNS[: data.shape[1]]
        return data

    out: dict[str, pd.Series] = {}
    for col in EXPECTED_COLUMNS:
        if col in headers:
            idx = headers.index(col)
            if idx < data.shape[1]:
                out[col] = data.iloc[:, idx]
            else:
                out[col] = pd.Series([None] * len(data))
        else:
            out[col] = pd.Series([None] * len(data))
    return pd.DataFrame(out)


def _read_org_excel(
    *,
    file_bytes: bytes | None = None,
    upload_filename: str | None = None,
    source_path: str | None = None,
) -> tuple[pd.DataFrame, str]:
    if file_bytes is not None:
        fn = (upload_filename or "upload.xlsx").lower()
        bio = io.BytesIO(file_bytes)
        engine = "xlrd" if fn.endswith(".xls") else "openpyxl"
        xls = pd.ExcelFile(bio, engine=engine)
        sheet = xls.sheet_names[0]
        raw = pd.read_excel(xls, sheet_name=sheet, header=None, dtype=str)
        return _frame_from_org_sheet(raw), sheet

    assert source_path
    raw = source_path.strip().strip('"').replace("\\", "/")
    p = Path(raw)
    if not p.is_absolute():
        p = (_project_root() / raw).resolve()
    if not p.is_file():
        raise FileNotFoundError(f"找不到文件：{p}")
    suf = p.suffix.lower()
    engine = "xlrd" if suf == ".xls" else "openpyxl"
    xls = pd.ExcelFile(p, engine=engine)
    sheet = xls.sheet_names[0]
    raw_df = pd.read_excel(p, sheet_name=sheet, header=None, dtype=str, engine=engine)
    return _frame_from_org_sheet(raw_df), sheet


def _display_name(shortname: str | None, fullname: str) -> str:
    return (shortname or "").strip() or fullname


def _compute_tree_fields(
    rows: list[dict[str, Any]],
    *,
    parent_key: str,
    sort_key: str,
    short_key: str,
    full_key: str,
) -> dict[str, dict[str, Any]]:
    by_id = {r["entity_id"]: r for r in rows}
    cache: dict[str, dict[str, Any]] = {}

    def walk(eid: str, seen: set[str]) -> dict[str, Any]:
        if eid in cache:
            return cache[eid]
        if eid in seen:
            raise ValueError(f"检测到 {parent_key} 环路：{eid}")
        row = by_id.get(eid)
        if not row:
            raise ValueError(f"父节点 {eid} 不存在于导入数据中")
        parent_id = str(row.get(parent_key) or "").strip()
        if parent_id == eid:
            level = 0
            path_names = [_display_name(row.get(short_key), row[full_key])]
            path_ids = [eid]
            root_group_id = None
        else:
            seen.add(eid)
            parent = walk(parent_id, seen)
            seen.discard(eid)
            level = int(parent["level"]) + 1
            path_names = list(parent["path_names"]) + [_display_name(row.get(short_key), row[full_key])]
            path_ids = list(parent["path_ids"]) + [eid]
            if level == 1:
                root_group_id = eid
            elif level == 0:
                root_group_id = None
            else:
                root_group_id = parent.get("root_group_id")
        children = [x["entity_id"] for x in rows if str(x.get(parent_key) or "").strip() == eid and x["entity_id"] != eid]
        result = {
            "level": level,
            "path_names": path_names,
            "path_ids_list": path_ids,
            "path": "/".join(path_names),
            "path_ids": "/".join(path_ids),
            "root_group_id": root_group_id,
            "is_leaf": len(children) == 0,
            "sort_no": row.get(sort_key),
        }
        cache[eid] = result
        return result

    for r in rows:
        walk(r["entity_id"], set())
    return cache


def _parse_rows_from_frame(df: pd.DataFrame, conn: Any) -> tuple[list[dict[str, Any]], list[str], list[str], list[dict[str, Any]]]:
    errors: list[str] = []
    warnings: list[str] = []
    reject_samples: list[dict[str, Any]] = []
    parsed: list[dict[str, Any]] = []

    try:
        sys_ids = {
            str(r[0]).strip()
            for r in conn.execute("SELECT sys_id FROM dim_org_sys WHERE is_active = TRUE").fetchall()
        }
    except Exception as exc:  # noqa: BLE001
        logger.warning("读取 dim_org_sys 失败: %s", exc)
        sys_ids = set()

    stat_years: set[int] = set()

    for idx, raw_row in df.iterrows():
        excel_row = int(idx) + 3
        if raw_row.isna().all():
            continue
        entity_id = _norm_cell(raw_row.get("entity_id"))
        if not entity_id:
            continue

        row_errors: list[str] = []

        entity_fullname = _norm_cell(raw_row.get("entity_fullname"))
        if not entity_fullname:
            row_errors.append(f"第 {excel_row} 行：entity_fullname 不能为空")

        entity_type = _norm_cell(raw_row.get("entity_type"))
        if not entity_type:
            row_errors.append(f"第 {excel_row} 行：entity_type 不能为空")

        sys_id = _norm_cell(raw_row.get("sys_id"))
        if not sys_id:
            row_errors.append(f"第 {excel_row} 行：sys_id 不能为空")
        elif sys_id not in sys_ids:
            row_errors.append(
                f"第 {excel_row} 行：sys_id={sys_id} 不在已启用的监管体系中，"
                "请先在「监管体系」维护或从模板「监管体系说明」页查询"
            )

        stat_year, err = _parse_int(raw_row.get("stat_year"), field="stat_year", row_no=excel_row, min_val=1990)
        if err:
            row_errors.append(err)
        elif stat_year and (stat_year < 1990 or stat_year > 2100):
            row_errors.append(f"第 {excel_row} 行：stat_year 超出有效范围")
        elif stat_year:
            stat_years.add(stat_year)

        is_stat_inc, err = _parse_bool_cn(raw_row.get("is_stat_inc"), field="is_stat_inc", row_no=excel_row)
        if err:
            row_errors.append(err)

        is_active, err = _parse_bool_cn(raw_row.get("is_active"), field="is_active", row_no=excel_row)
        if err:
            row_errors.append(err)

        mg_parent_id = _norm_cell(raw_row.get("mg_parent_id"))
        eq_parent_id = _norm_cell(raw_row.get("eq_parent_id"))
        if not mg_parent_id:
            row_errors.append(f"第 {excel_row} 行：mg_parent_id 不能为空")
        if not eq_parent_id:
            row_errors.append(f"第 {excel_row} 行：eq_parent_id 不能为空")

        mg_sort_no, err = _parse_int(raw_row.get("mg_sort_no"), field="mg_sort_no", row_no=excel_row, min_val=0)
        if err:
            row_errors.append(err)
        eq_sort_no, err = _parse_int(raw_row.get("eq_sort_no"), field="eq_sort_no", row_no=excel_row, min_val=0)
        if err:
            row_errors.append(err)

        eq_ratio, err = _parse_optional_decimal(
            raw_row.get("eq_shareholding_ratio"), field="eq_shareholding_ratio", row_no=excel_row
        )
        if err:
            row_errors.append(err)

        if not _valid_entity_id(entity_id):
            row_errors.append(f"第 {excel_row} 行：entity_id 须为 18 位统一社会信用代码或以 ROOT_ 开头")

        if mg_parent_id == entity_id and is_stat_inc is True:
            row_errors.append(f"第 {excel_row} 行：根节点（自引用）is_stat_inc 必须为「否」")

        if row_errors:
            errors.extend(row_errors)
            reject_samples.append(
                {
                    "seq_no": excel_row,
                    "sheet": "org_hier",
                    "field": "entity_id",
                    "reason": "；".join(row_errors[:3]),
                    "entity_id": entity_id,
                }
            )
            continue

        parsed.append(
            {
                "entity_id": entity_id,
                "entity_fullname": entity_fullname or "",
                "entity_shortname": _norm_cell(raw_row.get("entity_shortname")),
                "entity_type": entity_type or "",
                "sys_id": sys_id or "",
                "stat_year": int(stat_year or 0),
                "main_business": _norm_cell(raw_row.get("main_business")),
                "industry_id": _norm_cell(raw_row.get("industry_id")),
                "industry_name": _norm_cell(raw_row.get("industry_name")),
                "is_stat_inc": bool(is_stat_inc),
                "state_investor_id": _norm_cell(raw_row.get("state_investor_id")),
                "state_investor_name": _norm_cell(raw_row.get("state_investor_name")),
                "mg_parent_id": mg_parent_id or "",
                "mg_parent_name": _norm_cell(raw_row.get("mg_parent_name")),
                "mg_sort_no": int(mg_sort_no or 0),
                "eq_parent_id": eq_parent_id or "",
                "eq_parent_name": _norm_cell(raw_row.get("eq_parent_name")),
                "eq_sort_no": int(eq_sort_no or 0),
                "eq_shareholding_ratio": float(eq_ratio) if eq_ratio is not None else None,
                "reg_capital": _norm_cell(raw_row.get("reg_capital")),
                "is_active": bool(is_active),
                "hier_diff_note": _norm_cell(raw_row.get("hier_diff_note")),
                "excel_row": excel_row,
            }
        )

    if len(stat_years) > 1:
        errors.append(f"同一文件只允许一个 stat_year，当前发现：{sorted(stat_years)}")
    elif len(stat_years) == 0 and parsed:
        errors.append("未能解析 stat_year")

    entity_ids = {r["entity_id"] for r in parsed}
    by_id = {r["entity_id"]: r for r in parsed}
    for r in parsed:
        row_no = r["excel_row"]
        if r["mg_parent_id"] not in entity_ids:
            errors.append(f"第 {row_no} 行：mg_parent_id={r['mg_parent_id']} 不在本文件 entity_id 列中")
        if r["eq_parent_id"] not in entity_ids:
            errors.append(f"第 {row_no} 行：eq_parent_id={r['eq_parent_id']} 不在本文件 entity_id 列中")
        mg_parent = by_id.get(r["mg_parent_id"])
        if mg_parent and r.get("mg_parent_name"):
            expected = mg_parent["entity_fullname"]
            if r["mg_parent_name"] != expected:
                warnings.append(
                    f"第 {row_no} 行：mg_parent_name「{r['mg_parent_name']}」与上级 entity_fullname「{expected}」不一致"
                )
        eq_parent = by_id.get(r["eq_parent_id"])
        if eq_parent and r.get("eq_parent_name"):
            expected = eq_parent["entity_fullname"]
            if r["eq_parent_name"] != expected:
                warnings.append(
                    f"第 {row_no} 行：eq_parent_name「{r['eq_parent_name']}」与上级 entity_fullname「{expected}」不一致"
                )
        si_id = r.get("state_investor_id")
        if si_id and si_id not in entity_ids:
            errors.append(f"第 {row_no} 行：state_investor_id={si_id} 不在本文件 entity_id 列中")
        elif si_id:
            si_row = by_id.get(si_id)
            si_name = r.get("state_investor_name")
            if si_row and si_name and si_name != si_row["entity_fullname"]:
                warnings.append(
                    f"第 {row_no} 行：state_investor_name「{si_name}」与国家出资企业 entity_fullname「{si_row['entity_fullname']}」不一致"
                )

    if errors:
        return [], errors, warnings, reject_samples

    enriched, enrich_errors, enrich_warnings = _enrich_parsed_rows(parsed)
    warnings.extend(enrich_warnings)
    if enrich_errors:
        errors.extend(enrich_errors)
        return [], errors, warnings, reject_samples
    return enriched, errors, warnings, reject_samples


def _enrich_parsed_rows(parsed: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    if not parsed:
        return [], errors, warnings

    try:
        mg_fields = _compute_tree_fields(
            parsed,
            parent_key="mg_parent_id",
            sort_key="mg_sort_no",
            short_key="entity_shortname",
            full_key="entity_fullname",
        )
        eq_fields = _compute_tree_fields(
            parsed,
            parent_key="eq_parent_id",
            sort_key="eq_sort_no",
            short_key="entity_shortname",
            full_key="entity_fullname",
        )
    except ValueError as exc:
        errors.append(str(exc))
        return [], errors, warnings

    by_id = {r["entity_id"]: r for r in parsed}
    enriched: list[dict[str, Any]] = []
    for r in parsed:
        mg = mg_fields[r["entity_id"]]
        eq = eq_fields[r["entity_id"]]
        mg_parent = by_id.get(r["mg_parent_id"])
        eq_parent = by_id.get(r["eq_parent_id"])
        is_hier_diff = r["mg_parent_id"] != r["eq_parent_id"]
        row_no = r.get("excel_row", "?")
        if is_hier_diff and not r.get("hier_diff_note"):
            warnings.append(
                f"第 {row_no} 行：管产分离未填写 hier_diff_note（{r['entity_shortname'] or r['entity_fullname']}）"
            )

        si_id = (r.get("state_investor_id") or "").strip()
        si_name = (r.get("state_investor_name") or "").strip()
        mg_level = int(mg["level"])
        if not si_id:
            if mg_level == 1:
                si_id = r["entity_id"]
                si_name = si_name or r["entity_fullname"]
            elif mg_level >= 2:
                si_id = str(mg.get("root_group_id") or "").strip()
                if si_id:
                    anchor = by_id.get(si_id)
                    si_name = si_name or ((anchor or {}).get("entity_fullname") or "")
        elif mg_level == 1 and si_id != r["entity_id"]:
            warnings.append(
                f"第 {row_no} 行：一级企业国家出资企业代码一般填自身 entity_id（当前填 {si_id}）"
            )

        if r.get("is_stat_inc") and mg_level >= 1:
            if not si_id:
                errors.append(f"第 {row_no} 行：纳入统计企业须填写或可推断国家出资企业代码")
            elif si_id not in by_id:
                errors.append(f"第 {row_no} 行：state_investor_id={si_id} 不在本文件 entity_id 列中")
            else:
                si_anchor_level = int(mg_fields[si_id]["level"])
                if si_anchor_level != 1:
                    warnings.append(
                        f"第 {row_no} 行：国家出资企业「{si_id}」管理层级为 {si_anchor_level}，一般应为一级企业"
                    )
                si_row = by_id[si_id]
                expected_name = si_row["entity_fullname"]
                if si_name and si_name != expected_name:
                    warnings.append(
                        f"第 {row_no} 行：state_investor_name「{si_name}」与国家出资企业 entity_fullname「{expected_name}」不一致"
                    )
                elif not si_name:
                    si_name = expected_name

        enriched.append(
            {
                **r,
                "state_investor_id": si_id or None,
                "state_investor_name": si_name or None,
                "mg_level": mg_level,
                "mg_path": mg["path"],
                "mg_path_ids": mg["path_ids"],
                "mg_root_group_id": mg["root_group_id"],
                "mg_is_leaf": mg["is_leaf"],
                "eq_level": eq["level"],
                "eq_path": eq["path"],
                "eq_path_ids": eq["path_ids"],
                "eq_root_group_id": eq["root_group_id"],
                "eq_is_leaf": eq["is_leaf"],
                "mg_parent_shortname": (mg_parent or {}).get("entity_shortname"),
                "mg_parent_fullname": (mg_parent or {}).get("entity_fullname"),
                "eq_parent_shortname": (eq_parent or {}).get("entity_shortname"),
                "eq_parent_fullname": (eq_parent or {}).get("entity_fullname"),
                "is_hier_diff": is_hier_diff,
            }
        )

    return enriched, errors, warnings


def _write_yoy_logs(conn: Any, rows: list[dict[str, Any]], *, stat_year: int, changed_by: str = "SYSTEM") -> int:
    prior_year = stat_year - 1
    try:
        prior_rows = conn.execute(
            """
            SELECT entity_id, mg_parent_id, eq_parent_id
            FROM dim_org_hier
            WHERE stat_year = ?
            """,
            [prior_year],
        ).fetchall()
    except Exception as exc:  # noqa: BLE001
        logger.warning("读取上年度 dim_org_hier 失败: %s", exc)
        return 0

    prior_map = {str(r[0]): {"mg": r[1], "eq": r[2]} for r in prior_rows or []}
    log_count = 0
    for r in rows:
        eid = r["entity_id"]
        old = prior_map.get(eid)
        if not old:
            continue
        old_mg = old.get("mg")
        old_eq = old.get("eq")
        new_mg = r["mg_parent_id"]
        new_eq = r["eq_parent_id"]
        mg_changed = str(old_mg or "") != str(new_mg or "")
        eq_changed = str(old_eq or "") != str(new_eq or "")
        if not mg_changed and not eq_changed:
            continue
        change_type = "管理变更" if mg_changed and not eq_changed else "产权变更" if eq_changed and not mg_changed else "管产双变"
        note_parts = []
        if mg_changed:
            note_parts.append(f"管理上级 {old_mg} → {new_mg}")
        if eq_changed:
            note_parts.append(f"产权上级 {old_eq} → {new_eq}")
        conn.execute(
            """
            INSERT INTO dim_org_hier_log (
                entity_id, stat_year, change_type,
                old_mg_parent_id, new_mg_parent_id,
                old_eq_parent_id, new_eq_parent_id,
                change_note, changed_by
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                eid,
                stat_year,
                change_type,
                old_mg,
                new_mg,
                old_eq,
                new_eq,
                "；".join(note_parts),
                changed_by,
            ],
        )
        log_count += 1
    return log_count


def _persist_rows(conn: Any, rows: list[dict[str, Any]], *, stat_year: int, updated_by: str = "SYSTEM") -> tuple[int, int]:
    inserted = 0
    updated = 0
    for r in rows:
        conn.execute(
            """
            INSERT OR REPLACE INTO dim_org_node (
                entity_id, sys_id, entity_fullname, entity_shortname, entity_type,
                main_business, industry_id, industry_name, is_stat_inc, reg_capital,
                is_active, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            """,
            [
                r["entity_id"],
                r["sys_id"],
                r["entity_fullname"],
                r.get("entity_shortname"),
                r.get("entity_type"),
                r.get("main_business"),
                r.get("industry_id"),
                r.get("industry_name"),
                r.get("is_stat_inc", True),
                float(r["reg_capital"]) if r.get("reg_capital") not in (None, "") else None,
                r.get("is_active", True),
            ],
        )
        hid = _hier_id(r["entity_id"], stat_year)
        exists = conn.execute(
            "SELECT 1 FROM dim_org_hier WHERE entity_id = ? AND stat_year = ?",
            [r["entity_id"], stat_year],
        ).fetchone()
        conn.execute(
            "DELETE FROM dim_org_hier WHERE entity_id = ? AND stat_year = ?",
            [r["entity_id"], stat_year],
        )
        conn.execute(
            """
            INSERT INTO dim_org_hier (
                hier_id, entity_id, entity_shortname, entity_fullname, sys_id, stat_year,
                mg_parent_id, mg_parent_shortname, mg_parent_fullname, mg_sort_no,
                mg_level, mg_path, mg_path_ids, mg_root_group_id, mg_is_leaf,
                eq_parent_id, eq_parent_shortname, eq_parent_fullname, eq_sort_no,
                eq_shareholding_ratio, eq_level, eq_path, eq_path_ids, eq_root_group_id, eq_is_leaf,
                is_hier_diff, hier_diff_note, state_investor_id, state_investor_name, updated_at, updated_by
            ) VALUES (
                ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?, ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, CURRENT_TIMESTAMP, ?
            )
            """,
            [
                hid,
                r["entity_id"],
                r.get("entity_shortname"),
                r["entity_fullname"],
                r["sys_id"],
                stat_year,
                r["mg_parent_id"],
                r.get("mg_parent_shortname"),
                r.get("mg_parent_fullname"),
                r.get("mg_sort_no"),
                r.get("mg_level"),
                r.get("mg_path"),
                r.get("mg_path_ids"),
                r.get("mg_root_group_id"),
                r.get("mg_is_leaf"),
                r["eq_parent_id"],
                r.get("eq_parent_shortname"),
                r.get("eq_parent_fullname"),
                r.get("eq_sort_no"),
                r.get("eq_shareholding_ratio"),
                r.get("eq_level"),
                r.get("eq_path"),
                r.get("eq_path_ids"),
                r.get("eq_root_group_id"),
                r.get("eq_is_leaf"),
                r.get("is_hier_diff", False),
                r.get("hier_diff_note"),
                r.get("state_investor_id"),
                r.get("state_investor_name"),
                updated_by,
            ],
        )
        if exists:
            updated += 1
        else:
            inserted += 1
    return inserted, updated


def import_org_hierarchy_from_excel(
    conn: Any,
    *,
    file_bytes: bytes | None = None,
    upload_filename: str | None = None,
    source_path: str | None = None,
    dry_run: bool = False,
    run_id: str | None = None,
    on_progress: ProgressFn | None = None,
    updated_by: str = "SYSTEM",
) -> dict[str, Any]:
    def _prog(step: str, msg: str) -> None:
        if on_progress:
            on_progress(step, msg)

    build_run_id = run_id or f"org_hier_{uuid.uuid4().hex[:12]}"
    try:
        _prog("read_excel", "读取组织维度 Excel…")
        df, sheet = _read_org_excel(
            file_bytes=file_bytes,
            upload_filename=upload_filename,
            source_path=source_path,
        )
    except Exception as exc:  # noqa: BLE001
        logger.exception("读取组织 Excel 失败")
        return {
            "ok": False,
            "file_blocking": True,
            "error": {"message": str(exc), "exception_type": type(exc).__name__},
            "run_id": build_run_id,
        }

    _prog("validate", "校验模板字段…")
    rows, errors, warnings, reject_samples = _parse_rows_from_frame(df, conn)
    if errors:
        return {
            "ok": False,
            "file_blocking": False,
            "errors": errors,
            "warnings": warnings,
            "reject_row_samples": reject_samples,
            "run_id": build_run_id,
            "sheet": sheet,
        }

    stat_year = int(rows[0]["stat_year"]) if rows else 0
    hier_diff_count = sum(1 for r in rows if r.get("is_hier_diff"))

    if dry_run:
        return {
            "ok": True,
            "dry_run": True,
            "stat_year": stat_year,
            "row_count": len(rows),
            "hier_diff_count": hier_diff_count,
            "warnings": warnings,
            "reject_row_samples": reject_samples,
            "run_id": build_run_id,
            "sheet": sheet,
        }

    try:
        conn.execute("BEGIN")
        _prog("write_logs", "对比上年度并写入变更日志…")
        log_count = _write_yoy_logs(conn, rows, stat_year=stat_year, changed_by=updated_by)
        _prog("upsert", f"写入 dim_org_node / dim_org_hier（{stat_year}）…")
        inserted, updated = _persist_rows(conn, rows, stat_year=stat_year, updated_by=updated_by)
        conn.execute("COMMIT")
    except Exception as exc:  # noqa: BLE001
        try:
            conn.execute("ROLLBACK")
        except Exception:  # noqa: BLE001
            pass
        logger.exception("组织维度导入写入失败")
        return {
            "ok": False,
            "error": {"message": str(exc), "exception_type": type(exc).__name__},
            "run_id": build_run_id,
        }

    return {
        "ok": True,
        "dry_run": False,
        "stat_year": stat_year,
        "success": inserted + updated,
        "inserted": inserted,
        "updated": updated,
        "hier_diff_count": hier_diff_count,
        "log_count": log_count,
        "warnings": warnings,
        "reject_row_samples": reject_samples,
        "run_id": build_run_id,
        "sheet": sheet,
    }


def rebuild_org_hierarchy_paths(
    conn: Any,
    *,
    stat_years: list[int] | None = None,
    dry_run: bool = False,
    run_id: str | None = None,
) -> dict[str, Any]:
    """从 dim_org_hier 既有父节点关系重算路径/层级（不读 Excel）。"""
    build_run_id = run_id or f"org_hier_rebuild_{uuid.uuid4().hex[:12]}"
    try:
        if stat_years:
            years = [int(y) for y in stat_years]
        else:
            rows = conn.execute(
                "SELECT DISTINCT stat_year FROM dim_org_hier ORDER BY stat_year DESC"
            ).fetchall()
            years = [int(r[0]) for r in rows or []]
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "error": {"message": str(exc), "exception_type": type(exc).__name__}}

    if not years:
        return {
            "ok": False,
            "error": {
                "message": "dim_org_hier 中无数据，请先通过 Excel 导入组织层级模板",
                "exception_type": "ValidationError",
            },
            "run_id": build_run_id,
        }

    summaries: list[dict[str, Any]] = []
    for year in years:
        raw = conn.execute(
            """
            SELECT
                h.entity_id, h.sys_id, h.entity_fullname, h.entity_shortname,
                n.entity_type, n.main_business, n.industry_id, n.industry_name,
                COALESCE(n.is_stat_inc, TRUE), h.mg_parent_id, h.mg_sort_no,
                h.eq_parent_id, h.eq_sort_no, h.eq_shareholding_ratio,
                n.reg_capital, COALESCE(n.is_active, TRUE), h.hier_diff_note,
                h.state_investor_id, h.state_investor_name
            FROM dim_org_hier h
            LEFT JOIN dim_org_node n ON n.entity_id = h.entity_id
            WHERE h.stat_year = ?
            """,
            [year],
        ).fetchall()
        parsed = []
        for i, r in enumerate(raw or []):
            parsed.append(
                {
                    "entity_id": str(r[0]),
                    "sys_id": str(r[1] or ""),
                    "entity_fullname": str(r[2] or ""),
                    "entity_shortname": r[3],
                    "entity_type": str(r[4] or "企业"),
                    "main_business": r[5],
                    "industry_id": r[6],
                    "industry_name": r[7],
                    "is_stat_inc": bool(r[8]),
                    "mg_parent_id": str(r[9] or ""),
                    "mg_sort_no": int(r[10] or 0),
                    "eq_parent_id": str(r[11] or ""),
                    "eq_sort_no": int(r[12] or 0),
                    "eq_shareholding_ratio": float(r[13]) if r[13] is not None else None,
                    "reg_capital": r[14],
                    "is_active": bool(r[15]),
                    "hier_diff_note": r[16],
                    "state_investor_id": r[17],
                    "state_investor_name": r[18],
                    "excel_row": i + 1,
                }
            )
        entity_ids = {r["entity_id"] for r in parsed}
        errors = []
        for r in parsed:
            if r["mg_parent_id"] not in entity_ids:
                errors.append(f"entity_id={r['entity_id']} 的 mg_parent_id 不存在")
            if r["eq_parent_id"] not in entity_ids:
                errors.append(f"entity_id={r['entity_id']} 的 eq_parent_id 不存在")
        enriched, enrich_errors, warnings = _enrich_parsed_rows(parsed)
        errors.extend(enrich_errors)
        if errors:
            return {
                "ok": False,
                "stat_year": year,
                "errors": errors,
                "run_id": build_run_id,
            }
        if dry_run:
            summaries.append({"stat_year": year, "row_count": len(enriched), "dry_run": True})
            continue
        try:
            conn.execute("BEGIN")
            log_count = _write_yoy_logs(conn, enriched, stat_year=year)
            inserted, updated = _persist_rows(conn, enriched, stat_year=year)
            conn.execute("COMMIT")
        except Exception as exc:  # noqa: BLE001
            try:
                conn.execute("ROLLBACK")
            except Exception:  # noqa: BLE001
                pass
            return {"ok": False, "error": {"message": str(exc), "exception_type": type(exc).__name__}}
        summaries.append(
            {
                "stat_year": year,
                "row_count": len(enriched),
                "inserted": inserted,
                "updated": updated,
                "log_count": log_count,
                "warnings": warnings,
            }
        )

    return {
        "ok": True,
        "dry_run": dry_run,
        "run_id": build_run_id,
        "years": summaries,
        "rows_affected": sum(int(s.get("row_count") or 0) for s in summaries),
    }


def _parse_share_ratio(shareholders: str) -> float | None:
    from src.local_api.audited_enterprise_relation_api import _parse_first_shareholder

    _, ratio = _parse_first_shareholder(shareholders)
    if not ratio:
        return None
    s = ratio.strip().rstrip("%").strip()
    try:
        return float(s)
    except ValueError:
        return None


def materialize_org_hier_from_registry(
    conn: Any,
    *,
    stat_years: list[int] | None = None,
    replace_years: bool = True,
    run_id: str | None = None,
    on_progress: ProgressFn | None = None,
    updated_by: str = "REGISTRY_SYNC",
) -> dict[str, Any]:
    """
    从 dim_audited_enterprise_registry 物化 dim_org_node + dim_org_hier。

    台账为权威源：按年度覆盖写入，供 analysis_subject_pool 等仍读 dim_org_hier 的路径使用。
    """
    from src.local_api.group_enterprise_year_build import (
        _norm_id,
        _norm_name,
        _parse_equity_parent_name,
        _registry_years,
        _resolve_parent_id,
        _safe_level,
    )

    def _prog(step: str, msg: str) -> None:
        if on_progress:
            on_progress(step, msg)

    years = [int(y) for y in (stat_years or []) if y is not None]
    if not years:
        years = _registry_years(conn)
    if not years:
        return {
            "ok": False,
            "error": {
                "message": "dim_audited_enterprise_registry 无可用 snapshot_year",
                "exception_type": "ValidationError",
            },
        }

    sys_ids = _load_active_sys_ids(conn)
    default_sys_id = sys_ids[0] if sys_ids else "DEFAULT"
    build_run_id = run_id or f"org_hier_mat_{uuid.uuid4().hex[:12]}"
    summaries: list[dict[str, Any]] = []
    total_written = 0

    for year_i in years:
        _prog("load_registry", f"读取 {year_i} 年度台账以物化 dim_org_hier…")
        try:
            rows = conn.execute(
                """
                SELECT
                    unified_social_credit_code,
                    enterprise_name,
                    main_business,
                    enterprise_category,
                    registered_capital,
                    state_investor,
                    mgmt_level,
                    mgmt_parent,
                    equity_level,
                    shareholders
                FROM dim_audited_enterprise_registry
                WHERE snapshot_year = ?
                  AND trim(COALESCE(unified_social_credit_code, '')) <> ''
                """,
                [year_i],
            ).fetchall()
        except Exception as exc:  # noqa: BLE001
            logger.exception("读取台账失败 year=%s", year_i)
            return {"ok": False, "error": {"message": str(exc), "exception_type": type(exc).__name__}}

        if not rows:
            summaries.append({"stat_year": year_i, "written": 0, "skipped": True, "reason": "台账无行"})
            continue

        name_to_id: dict[str, str] = {}
        id_to_name: dict[str, str] = {}
        parsed_rows: list[dict[str, Any]] = []
        for r in rows or []:
            eid = _norm_id(str(r[0] or ""))
            if not eid or not _valid_entity_id(eid):
                continue
            ename = str(r[1] or "").strip() or eid
            parsed_rows.append(
                {
                    "enterprise_id": eid,
                    "enterprise_name": ename,
                    "main_business": str(r[2] or "").strip(),
                    "enterprise_category": str(r[3] or "").strip() or "企业",
                    "registered_capital": str(r[4] or "").strip(),
                    "state_investor": str(r[5] or "").strip(),
                    "mgmt_level": _safe_level(r[6], 2),
                    "mgmt_parent_name": str(r[7] or "").strip(),
                    "equity_level": _safe_level(r[8], 2),
                    "shareholders": str(r[9] or "").strip(),
                }
            )
            id_to_name[eid] = ename
            nk = _norm_name(ename)
            if nk and nk not in name_to_id:
                name_to_id[nk] = eid

        parsed: list[dict[str, Any]] = []
        for row in parsed_rows:
            eid = row["enterprise_id"]
            mgmt_level = int(row["mgmt_level"])
            equity_level = int(row["equity_level"])
            mgmt_parent_name = row["mgmt_parent_name"]
            equity_parent_name = _parse_equity_parent_name(row["shareholders"])

            mg_parent_id, _ = _resolve_parent_id(
                mgmt_parent_name,
                self_id=eid,
                level=mgmt_level,
                name_to_id=name_to_id,
            )
            eq_parent_id, _ = _resolve_parent_id(
                equity_parent_name,
                self_id=eid,
                level=equity_level,
                name_to_id=name_to_id,
            )

            si_name = row["state_investor"]
            si_id = ""
            if si_name:
                si_id = name_to_id.get(_norm_name(si_name), "")
            if mgmt_level <= 1 and not si_id:
                si_id = eid
                si_name = si_name or row["enterprise_name"]

            eq_ratio = _parse_share_ratio(row["shareholders"])

            parsed.append(
                {
                    "entity_id": eid,
                    "entity_fullname": row["enterprise_name"],
                    "entity_shortname": row["enterprise_name"][:32] if row["enterprise_name"] else eid,
                    "entity_type": row["enterprise_category"] or "企业",
                    "sys_id": default_sys_id,
                    "stat_year": year_i,
                    "main_business": row["main_business"] or None,
                    "industry_id": None,
                    "industry_name": None,
                    "is_stat_inc": True,
                    "state_investor_id": si_id or None,
                    "state_investor_name": si_name or None,
                    "mg_parent_id": mg_parent_id,
                    "mg_parent_name": mgmt_parent_name or id_to_name.get(mg_parent_id, ""),
                    "mg_sort_no": mgmt_level,
                    "eq_parent_id": eq_parent_id,
                    "eq_parent_name": equity_parent_name or id_to_name.get(eq_parent_id, ""),
                    "eq_sort_no": equity_level,
                    "eq_shareholding_ratio": eq_ratio,
                    "reg_capital": row["registered_capital"] or None,
                    "is_active": True,
                    "hier_diff_note": None,
                    "excel_row": 0,
                }
            )

        entity_ids = {r["entity_id"] for r in parsed}
        errors: list[str] = []
        for r in parsed:
            if r["mg_parent_id"] not in entity_ids:
                errors.append(f"{r['entity_id']} 的管理上级未在台账中匹配")
            if r["eq_parent_id"] not in entity_ids:
                errors.append(f"{r['entity_id']} 的产权上级未在台账中匹配")
        if errors:
            summaries.append(
                {
                    "stat_year": year_i,
                    "written": 0,
                    "skipped": True,
                    "reason": errors[0],
                    "errors": errors[:5],
                }
            )
            continue

        enriched, enrich_errors, warnings = _enrich_parsed_rows(parsed)
        if enrich_errors:
            summaries.append(
                {
                    "stat_year": year_i,
                    "written": 0,
                    "skipped": True,
                    "reason": enrich_errors[0],
                    "errors": enrich_errors[:5],
                }
            )
            continue

        _prog("write_org_hier", f"写入 {year_i} 年度 dim_org_hier（{len(enriched)} 行）…")
        try:
            conn.execute("BEGIN")
            if replace_years:
                conn.execute("DELETE FROM dim_org_hier WHERE stat_year = ?", [year_i])
            log_count = _write_yoy_logs(conn, enriched, stat_year=year_i, changed_by=updated_by)
            inserted, updated = _persist_rows(conn, enriched, stat_year=year_i, updated_by=updated_by)
            conn.execute("COMMIT")
        except Exception as exc:  # noqa: BLE001
            try:
                conn.execute("ROLLBACK")
            except Exception:  # noqa: BLE001
                pass
            logger.exception("物化 dim_org_hier 失败 year=%s", year_i)
            return {"ok": False, "error": {"message": str(exc), "exception_type": type(exc).__name__}}

        total_written += len(enriched)
        hier_diff_count = sum(1 for x in enriched if x.get("is_hier_diff"))
        summaries.append(
            {
                "stat_year": year_i,
                "written": len(enriched),
                "inserted": inserted,
                "updated": updated,
                "log_count": log_count,
                "hier_diff_count": hier_diff_count,
                "warnings": warnings[:10],
            }
        )

    if total_written == 0 and any(s.get("skipped") for s in summaries):
        first_err = next((s.get("reason") for s in summaries if s.get("skipped")), "物化失败")
        return {
            "ok": False,
            "error": {"message": str(first_err), "exception_type": "MaterializeError"},
            "run_id": build_run_id,
            "year_summaries": summaries,
        }

    return {
        "ok": True,
        "run_id": build_run_id,
        "rows_written": total_written,
        "stat_years": years,
        "year_summaries": summaries,
    }
