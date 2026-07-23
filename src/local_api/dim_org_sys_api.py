"""监管体系 dim_org_sys CRUD API。"""

from __future__ import annotations

import io
import logging
import re
from typing import Any

logger = logging.getLogger(__name__)

_SYS_ID_RE = re.compile(r"^[A-Z][A-Z0-9_]{1,31}$")
_ADMIN_LEVELS = ("省", "市", "区/县", "中央", "其他")
_EXPORT_FILENAME = "监管体系清单.xlsx"

# (sys_id, sys_name, admin_level, gov_owner, description, sort_no, is_active)
_DEFAULT_TEMPLATES: tuple[tuple[Any, ...], ...] = (
    (
        "PROV_XX",
        "省属企业",
        "省",
        "XX省国有资产监督管理委员会",
        "如：山东省属国有企业，由省国资委统一监管",
        10,
        False,
    ),
    (
        "CITY_XX",
        "市属企业",
        "市",
        "XX市国有资产监督管理委员会",
        "如：济南市属国有企业，由市国资委监管",
        20,
        False,
    ),
    (
        "DIST_XX",
        "区县属企业",
        "区/县",
        "XX区（县）国有资产监督管理机构",
        "如：某区（县）属国有企业",
        25,
        False,
    ),
    (
        "CENTRAL",
        "中央企业",
        "中央",
        "国务院国有资产监督管理委员会",
        "国务院国资委监管的中央企业体系",
        30,
        False,
    ),
    (
        "CULT_XX",
        "文化类企业",
        "省",
        "XX省文化和旅游厅",
        "如：山东省文化类国有企业",
        40,
        False,
    ),
    (
        "FIN_XX",
        "金融类国资",
        "省",
        "XX省财政厅",
        "如：山东省财政厅分管的金融类国有企业；省属金融类亦可能由省国资委监管",
        50,
        False,
    ),
)

_EXPORT_COL_WIDTHS = (16, 14, 10, 22, 40, 8, 8)


def _norm_str(v: Any) -> str:
    return str(v or "").strip()


def _load_scope() -> dict[str, Any]:
    try:
        from config.settings import SCOPE

        return dict(SCOPE) if isinstance(SCOPE, dict) else {}
    except Exception:  # noqa: BLE001
        return {}


def _virtual_root_entity_id(sys_id: str) -> str:
    """组织树虚拟根节点 entity_id，约定为 ROOT_{sys_id}。"""
    return f"ROOT_{sys_id}"


def _ref_counts(conn: Any, sys_id: str) -> tuple[int, int]:
    """统计可被组织数据引用的节点/层级条数（不含虚拟根 ROOT_{sys_id}）。"""
    root_id = _virtual_root_entity_id(sys_id)
    node_cnt = conn.execute(
        "SELECT COUNT(*) FROM dim_org_node WHERE sys_id = ? AND entity_id <> ?",
        [sys_id, root_id],
    ).fetchone()[0]
    hier_cnt = conn.execute(
        "SELECT COUNT(*) FROM dim_org_hier WHERE sys_id = ? AND entity_id <> ?",
        [sys_id, root_id],
    ).fetchone()[0]
    return int(node_cnt or 0), int(hier_cnt or 0)


def _row_to_item(row: tuple[Any, ...]) -> dict[str, Any]:
    return {
        "sys_id": _norm_str(row[0]),
        "sys_name": _norm_str(row[1]),
        "admin_level": _norm_str(row[2]),
        "gov_owner": _norm_str(row[3]),
        "description": _norm_str(row[4]),
        "sort_no": int(row[5] or 0),
        "is_active": bool(row[6]) if row[6] is not None else True,
        "node_count": int(row[7] or 0),
        "hier_count": int(row[8] or 0),
    }


def _list_items(conn: Any, *, active_only: bool = False) -> list[dict[str, Any]]:
    sql = """
        SELECT
            s.sys_id,
            s.sys_name,
            s.admin_level,
            COALESCE(s.gov_owner, ''),
            COALESCE(s.description, ''),
            s.sort_no,
            s.is_active,
            (
                SELECT COUNT(*) FROM dim_org_node n
                WHERE n.sys_id = s.sys_id AND n.entity_id <> 'ROOT_' || s.sys_id
            ) AS node_count,
            (
                SELECT COUNT(*) FROM dim_org_hier h
                WHERE h.sys_id = s.sys_id AND h.entity_id <> 'ROOT_' || s.sys_id
            ) AS hier_count
        FROM dim_org_sys s
    """
    if active_only:
        sql += " WHERE COALESCE(s.is_active, TRUE)"
    sql += " ORDER BY s.sort_no, s.sys_id"
    rows = conn.execute(sql).fetchall()
    return [_row_to_item(r) for r in rows or []]


def _seed_missing_templates(conn: Any) -> int:
    """补全缺失的初始模板（均为停用，不覆盖已有记录）。返回本次新增条数。"""
    inserted = 0
    try:
        for tpl in _DEFAULT_TEMPLATES:
            hit = conn.execute(
                "SELECT 1 FROM dim_org_sys WHERE sys_id = ? LIMIT 1",
                [tpl[0]],
            ).fetchone()
            if hit is not None:
                continue
            conn.execute(
                """
                INSERT INTO dim_org_sys (
                    sys_id, sys_name, admin_level, gov_owner, description, sort_no, is_active
                )
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                list(tpl),
            )
            inserted += 1
    except Exception as exc:  # noqa: BLE001
        logger.warning("org_sys seed templates: %s", exc)
    return inserted


def api_org_sys_list(conn: Any, *, active_only: bool = False) -> dict[str, Any]:
    try:
        items = _list_items(conn, active_only=active_only)
    except Exception as exc:  # noqa: BLE001
        logger.warning("org_sys list: %s", exc)
        return {"ok": False, "error": {"message": str(exc), "exception_type": type(exc).__name__}}
    scope = _load_scope()
    scope_sys_id = _norm_str(scope.get("scope_sys_id"))
    empty_hint = ""
    if not items:
        empty_hint = (
            "尚未维护任何监管体系。可点击「加载初始模板」补全示例 sys_id（按代码精确匹配、不覆盖已有记录），"
            "或手动新增（如 PROV_SD 山东省属企业），再下载组织维度导入模板。"
        )
    return {
        "ok": True,
        "items": items,
        "scope_sys_id": scope_sys_id,
        "scope_root_id": _norm_str(scope.get("scope_root_id")),
        "instance_name": _norm_str(scope.get("instance_name")),
        "empty_hint": empty_hint,
        "admin_levels": list(_ADMIN_LEVELS),
    }


def _validate_item(body: dict[str, Any], *, is_create: bool) -> tuple[dict[str, Any] | None, str | None]:
    sys_id = _norm_str(body.get("sys_id")).upper()
    sys_name = _norm_str(body.get("sys_name"))
    admin_level = _norm_str(body.get("admin_level"))
    gov_owner = _norm_str(body.get("gov_owner"))
    description = _norm_str(body.get("description"))
    sort_raw = body.get("sort_no")
    is_active = body.get("is_active")
    if is_create and not sys_id:
        return None, "sys_id 不能为空"
    if sys_id and not _SYS_ID_RE.match(sys_id):
        return None, "sys_id 须为大写字母/数字/下划线，如 PROV_SD、CITY_JN"
    if not sys_name:
        return None, "sys_name（名称）不能为空"
    if admin_level not in _ADMIN_LEVELS:
        return None, f"admin_level 须为：{' / '.join(_ADMIN_LEVELS)}"
    sort_no = 0
    if sort_raw is not None and str(sort_raw).strip() != "":
        try:
            sort_no = int(sort_raw)
        except (TypeError, ValueError):
            return None, "sort_no 须为整数"
    active = True if is_active is None else bool(is_active)
    return {
        "sys_id": sys_id,
        "sys_name": sys_name,
        "admin_level": admin_level,
        "gov_owner": gov_owner or None,
        "description": description or None,
        "sort_no": sort_no,
        "is_active": active,
    }, None


def api_org_sys_save(conn: Any, body: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(body, dict):
        return {"ok": False, "error": {"message": "请求体须为 JSON 对象"}}
    sys_id_raw = _norm_str(body.get("sys_id")).upper()
    exists = False
    if sys_id_raw:
        try:
            hit = conn.execute(
                "SELECT 1 FROM dim_org_sys WHERE sys_id = ? LIMIT 1",
                [sys_id_raw],
            ).fetchone()
            exists = hit is not None
        except Exception as exc:  # noqa: BLE001
            return {"ok": False, "error": {"message": str(exc), "exception_type": type(exc).__name__}}
    item, err = _validate_item(body, is_create=not exists)
    if err or not item:
        return {"ok": False, "error": {"message": err or "校验失败", "exception_type": "ValidationError"}}
    try:
        conn.execute(
            """
            INSERT INTO dim_org_sys (
                sys_id, sys_name, admin_level, gov_owner, description, sort_no, is_active
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT (sys_id) DO UPDATE SET
                sys_name = excluded.sys_name,
                admin_level = excluded.admin_level,
                gov_owner = excluded.gov_owner,
                description = excluded.description,
                sort_no = excluded.sort_no,
                is_active = excluded.is_active
            """,
            [
                item["sys_id"],
                item["sys_name"],
                item["admin_level"],
                item["gov_owner"],
                item["description"],
                item["sort_no"],
                item["is_active"],
            ],
        )
    except Exception as exc:  # noqa: BLE001
        logger.warning("org_sys save: %s", exc)
        return {"ok": False, "error": {"message": str(exc), "exception_type": type(exc).__name__}}
    return {"ok": True, "item": item, "created": not exists}


def api_org_sys_delete(conn: Any, body: dict[str, Any]) -> dict[str, Any]:
    sys_id = _norm_str((body or {}).get("sys_id")).upper()
    if not sys_id:
        return {"ok": False, "error": {"message": "sys_id 不能为空", "exception_type": "ValidationError"}}
    try:
        exists = conn.execute(
            "SELECT 1 FROM dim_org_sys WHERE sys_id = ? LIMIT 1",
            [sys_id],
        ).fetchone()
        if exists is None:
            return {"ok": False, "error": {"message": f"sys_id={sys_id} 不存在", "exception_type": "NotFoundError"}}
        node_cnt, hier_cnt = _ref_counts(conn, sys_id)
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "error": {"message": str(exc), "exception_type": type(exc).__name__}}
    ref_total = node_cnt + hier_cnt
    root_id = _virtual_root_entity_id(sys_id)
    try:
        if ref_total > 0:
            conn.execute(
                "UPDATE dim_org_sys SET is_active = FALSE WHERE sys_id = ?",
                [sys_id],
            )
            return {
                "ok": True,
                "deactivated": True,
                "message": f"sys_id={sys_id} 已被 {ref_total} 条组织数据引用，已改为停用而非删除",
            }
        conn.execute("DELETE FROM dim_org_hier WHERE sys_id = ? AND entity_id = ?", [sys_id, root_id])
        conn.execute("DELETE FROM dim_org_node WHERE sys_id = ? AND entity_id = ?", [sys_id, root_id])
        conn.execute("DELETE FROM dim_org_sys WHERE sys_id = ?", [sys_id])
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "error": {"message": str(exc), "exception_type": type(exc).__name__}}
    return {"ok": True, "deleted": True}


def api_org_sys_bootstrap_demo(conn: Any) -> dict[str, Any]:
    """补全缺失的初始监管体系模板（均为停用，不覆盖已有记录）。"""
    try:
        inserted = _seed_missing_templates(conn)
        items = _list_items(conn)
        message = (
            f"已加载 {inserted} 条缺失的初始模板（均为停用），请按需修改并启用"
            if inserted > 0
            else "初始模板 sys_id 均已存在，未新增记录"
        )
        return {"ok": True, "items": items, "seeded_count": inserted, "message": message}
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "error": {"message": str(exc), "exception_type": type(exc).__name__}}


def api_org_sys_export(conn: Any) -> tuple[int, bytes | dict[str, Any], str, str]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        items = _list_items(conn)
        wb = Workbook()
        ws = wb.active
        ws.title = "监管体系说明"
        headers = [
            "sys_id（监管体系代码）",
            "sys_name（名称）",
            "admin_level（层级）",
            "gov_owner（主管单位）",
            "description（说明）",
            "is_active（启用）",
        ]
        ws.append(headers)
        bold = Font(bold=True)
        for cell in ws[1]:
            cell.font = bold
        for row in items:
            ws.append(
                [
                    row["sys_id"],
                    row["sys_name"],
                    row["admin_level"],
                    row["gov_owner"],
                    row["description"],
                    "是" if row["is_active"] else "否",
                ]
            )
        ws.append([])
        ws.append(["说明", "sys_id 为系统预置字典项，导入组织维度时须从本清单选择，不可自由填写。"])
        ws.append(["维护入口", "维度管理 → 企业组织维度 → 监管体系"])
        if not items:
            ws.append(["提示", "当前清单为空，请先在系统中维护监管体系后再导入组织维度。"])
        for col_idx, width in enumerate(_EXPORT_COL_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        bio = io.BytesIO()
        wb.save(bio)
        return 200, bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _EXPORT_FILENAME
    except Exception as exc:  # noqa: BLE001
        logger.warning("org_sys export: %s", exc)
        return 500, {"ok": False, "error": {"message": str(exc), "exception_type": type(exc).__name__}}, "application/json", ""
