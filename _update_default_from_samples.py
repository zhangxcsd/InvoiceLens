# -*- coding: utf-8 -*-
from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

ROOT = Path(r"d:\\PythonCode\\InvoiceLens")
DATA = ROOT / "Source_Data" / "Invoice"

FIELD_KEYS = [
    "invoice_no",
    "invoice_code",
    "sdfphm",
    "kprq",
    "seller_tax_no",
    "buyer_tax_no",
    "amount",
    "tax_amount",
    "total_amount",
]


def norm(s: Any) -> str:
    if s is None:
        return ""
    return str(s).strip().replace(" ", "")


def load_sheet_keys(path: Path) -> dict[str, str]:
    raw = path.read_text(encoding="utf-8")
    data = yaml.safe_load(raw)
    if not isinstance(data, dict):
        return {}
    return {str(k): str(v) for k, v in data.items() if isinstance(k, str)}


def sheet_key_for_name(sheet_name: str, sheet_mapping: dict[str, str]) -> str | None:
    n = norm(sheet_name)
    best = None
    best_len = -1
    for k in sheet_mapping.keys():
        kn = norm(k)
        if kn and kn in n and len(kn) > best_len:
            best = k
            best_len = len(kn)
    return best


def detect_header_row(ws, max_scan_rows: int = 12, min_nonempty: int = 5) -> list[str] | None:
    max_r = min(max_scan_rows, int(ws.max_row or 0))
    max_c = int(ws.max_column or 0)
    if max_r <= 0 or max_c <= 0:
        return None
    best: tuple[int, list[str]] | None = None
    for r in range(1, max_r + 1):
        texts: list[str] = []
        nonempty = 0
        for c in range(1, max_c + 1):
            v = ws.cell(row=r, column=c).value
            t = "" if v is None else str(v).strip()
            texts.append(t)
            if t:
                nonempty += 1
        if nonempty >= min_nonempty:
            while texts and not texts[-1]:
                texts.pop()
            score = (nonempty, len(texts))
            if best is None or score > (best[0], len(best[1])):
                best = (nonempty, texts)
    return best[1] if best else None


def classify_header_to_fields(header: str) -> set[str]:
    n = norm(header)
    if not n:
        return set()
    out: set[str] = set()
    if "数电" in n and ("号码" in n or "票号" in n):
        out.add("sdfphm")
        return out
    if "发票代码" in n or n == "代码":
        out.add("invoice_code")
    if "发票号码" in n or ("发票号" in n and "代码" not in n):
        out.add("invoice_no")
    if "开票日期" in n or n == "日期":
        out.add("kprq")
    seller_hit = ("销方" in n) or ("销售方" in n) or ("销货方" in n)
    if seller_hit and ("税号" in n or "识别号" in n or "信用代码" in n or n.endswith("识别号")):
        out.add("seller_tax_no")
    buyer_hit = ("购方" in n) or ("购买方" in n) or ("购货方" in n)
    if buyer_hit and ("税号" in n or "识别号" in n or "信用代码" in n or n.endswith("识别号")):
        out.add("buyer_tax_no")
    if n in ("金额", "不含税金额"):
        out.add("amount")
    if n == "税额" or (n.startswith("税额") and "合计" not in n):
        out.add("tax_amount")
    if "价税合计" in n or n == "含税金额":
        out.add("total_amount")
    return out


sheet_mapping = load_sheet_keys(ROOT / "config" / "sheet_mapping.yaml")
TARGET_SHEETS = {"发票基础信息", "信息汇总表"}

aliases: dict[str, set[str]] = {k: set() for k in FIELD_KEYS}

for path in sorted(DATA.rglob("*.xlsx")):
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        continue
    try:
        for sn in wb.sheetnames:
            sk = sheet_key_for_name(sn, sheet_mapping)
            if sk not in TARGET_SHEETS:
                continue
            ws = wb[sn]
            header = detect_header_row(ws)
            if not header:
                continue
            for h in header:
                if not str(h).strip():
                    continue
                hs = str(h).strip()
                for f in classify_header_to_fields(hs):
                    if f in aliases:
                        aliases[f].add(hs)
    finally:
        wb.close()

cfg_path = ROOT / "config" / "field_mapping.yaml"
raw = cfg_path.read_text(encoding="utf-8")
doc = yaml.safe_load(raw) or {}
if "default" not in doc or not isinstance(doc["default"], dict):
    doc["default"] = {}

for f in FIELD_KEYS:
    vals = sorted(aliases[f], key=lambda x: (len(x), x))
    if not vals:
        continue
    doc["default"][f] = vals

cfg_path.write_text(
    "# 根据样例中“发票基础信息”“信息汇总表”两类 sheet 归纳 default 别名\n"
    + yaml.safe_dump(doc, allow_unicode=True, sort_keys=False, default_flow_style=False),
    encoding="utf-8",
)

print({k: sorted(v) for k, v in aliases.items()})
