# -*- coding: utf-8 -*-
"""
从样例 xlsx 中读取工作表名包含指定 key 的 sheet 表头，更新 config/field_mapping.yaml 中 sheets.<sheet_key> 块。

default 与 sheets 的关系：
- default 仅表示「多 sheet 共享」的公共别名子集，与各 sheets.* 块不是超集/子集包含关系。
- 表头匹配顺序：① default 中别名（按 default 字段顺序先匹配先赢）② 当前 YAML 里该 sheet 块已有别名
  ③ 仍无法匹配：写入本 sheet 专属字段 key（来自 YAML 顶层 sheet_header_slugs；未配置则自动生成可读 slug 并写回 YAML），无需事先在 default 注册。

其它规则：
- 列名归一化：strip + 去掉半角空格（与 excel_to_ods 一致）
- 每个字段最终别名集合 =（若 default 有该 key）default 别名 ∪ 观测列名；（sheet 专属 key 无 default 时仅为观测列名）
- 若合并结果与 default 中该字段别名集合完全相同，YAML 写 []；否则写排序后的列表
- 默认只写入本 sheet 样例中出现过的字段 key；需要「列满 default 全部 key」时加 --emit-all-default-keys
- 无样例时：9 个核心字段各 [] + 保留该 sheet 下已有非空手工别名

用法（项目根）：
  python tools/sync_sheet_block_from_samples.py --sheet-key 发票基础信息 --match contains
  python tools/sync_sheet_block_from_samples.py --all-sheet-keys --match contains
  # 合并多个样例根目录（先扫完第一个目录再扫第二个，列观测取并集）
  python tools/sync_sheet_block_from_samples.py --data Source_Data/Invoice --data D:/path/to/more --all-sheet-keys
  # 试跑：只处理排序后的前 N 个 xlsx（0 表示不限制）
  python tools/sync_sheet_block_from_samples.py --all-sheet-keys --max-excel-files 200
"""
from __future__ import annotations

import argparse
import warnings
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterator

import yaml
from openpyxl import load_workbook

from src.services.field_slug import suggest_field_slug

warnings.filterwarnings(
    "ignore",
    message="Workbook contains no default style",
    category=UserWarning,
    module="openpyxl.styles.stylesheet",
)

ROOT = Path(__file__).resolve().parents[1]
DATA_DEFAULT = ROOT / "Source_Data" / "Invoice"
CFG_PATH = ROOT / "config" / "field_mapping.yaml"
SHEET_MAP_PATH = ROOT / "config" / "sheet_mapping.yaml"

# 样例中未出现该 sheet 时，占位用的「通用字段」（与 ODS 核心列一致，便于勾选导入）
DEFAULT_CORE_FIELDS: list[str] = [
    "invoice_no",
    "invoice_code",
    "sdfphm",
    "kprq",
    "seller_tax_no",
    "buyer_tax_no",
    "amount",
    "tax_amount",
    "total_amount",
]

# 仅用于“脚本首次自动整理”时给出更可读的默认建议（最终会写入 YAML 的 sheet_header_slugs）。
# 该表不参与运行时推断；若你想完全由 YAML 驱动，可把 YAML 里维护好的映射固定后，此表随时可删。
SUGGESTED_SHEET_HEADER_SLUGS: dict[str, str] = {
    "车辆识别代号/车架号": "vin_chassis",
    "合格证号": "vehicle_certificate_no",
    "发动机号码": "engine_no",
    "商检单号": "commodity_inspection_no",
    "商检号码": "commodity_inspection_no",
    "产地": "place_of_origin",
    "跨区域涉税事项报验管理编号": "cross_region_tax_mgmt_no",
    "建筑服务发生地": "construction_service_location",
    "建筑项目名称": "construction_project_name",
    "产权证书/不动产权证号": "property_title_cert_no",
    "车牌号": "license_plate_no",
    "运输工具种类": "transport_means_type",
    "运输工具牌号": "transport_means_plate_no",
    "起运地": "shipping_origin",
    "到达地": "arrival_place",
    "运输货物名称": "cargo_name",
    "出行人": "trip_traveler",
    "有效身份证号": "traveler_id_no",
    "出行日期": "trip_date",
    "出发地": "departure_place",
    "等级": "service_class",
    "交通工具类型": "transport_tool_type",
    "旅客姓名": "passenger_name",
    "出行车次": "train_or_flight_no",
    "日期": "rail_journey_date",
    "时间": "rail_journey_time",
    # 二手车销售等
    "销方地址": "seller_address",
    "销货方联系电话": "seller_phone",
    "购买方地址": "buyer_address",
    "购买方联系电话": "buyer_phone",
    "车牌照号": "vehicle_plate_no",
    "登记证号": "registration_cert_no",
    "转入地车辆管理所名称": "transfer_dmv_office",
    "经营/拍卖单位": "auction_operator",
    "经营/拍卖单位纳税人识别号": "auction_operator_tax_no",
    "二手车市场": "used_vehicle_market",
    "二手车市场纳税人识别号": "used_vehicle_market_tax_no",
}


def _norm(s: Any) -> str:
    if s is None:
        return ""
    return str(s).strip().replace(" ", "")


def _detect_header_row(ws: Any, *, max_scan_rows: int = 12, min_nonempty: int = 5) -> list[str] | None:
    max_r = min(max_scan_rows, int(ws.max_row or 0))
    max_c = int(ws.max_column or 0)
    if max_r <= 0 or max_c <= 0:
        return None
    best: tuple[int, list[str]] | None = None
    for r in range(1, max_r + 1):
        texts: list[str] = []
        nonempty = 0
        for c in range(1, max_c + 1):
            v = ws.cell(row=r, column=c).value
            t = "" if v is None else str(v).strip()
            texts.append(t)
            if t:
                nonempty += 1
        if nonempty >= min_nonempty:
            while texts and not texts[-1]:
                texts.pop()
            score = (nonempty, len(texts))
            if best is None or score > (best[0], len(best[1])):
                best = (nonempty, texts)
    return best[1] if best else None


def _default_field_order_and_aliases(doc: dict[str, Any]) -> tuple[list[str], dict[str, list[str]]]:
    d = doc.get("default")
    if not isinstance(d, dict):
        return [], {}
    order: list[str] = []
    aliases: dict[str, list[str]] = {}
    for k, v in d.items():
        if not isinstance(k, str):
            continue
        key = k.strip()
        if not key:
            continue
        if not isinstance(v, list):
            continue
        als = [str(x).strip() for x in v if x is not None and str(x).strip()]
        order.append(key)
        aliases[key] = als
    return order, aliases


def _load_sheet_mapping_keys(path: Path) -> list[str]:
    raw = path.read_text(encoding="utf-8")
    data = yaml.safe_load(raw)
    if not isinstance(data, dict):
        return []
    return [str(k).strip() for k in data.keys() if isinstance(k, str) and str(k).strip()]


def _header_to_field(header: str, field_order: list[str], default_aliases: dict[str, list[str]]) -> str | None:
    hn = _norm(header)
    if not hn:
        return None
    for field in field_order:
        for a in default_aliases.get(field, []):
            if _norm(a) == hn:
                return field
    return None


def _ensure_sheet_header_slugs(doc: dict[str, Any]) -> dict[str, str]:
    """
    返回 YAML 顶层 sheet_header_slugs（按“原始表头文本”作为 key）。
    若不存在会创建为空 dict。
    """
    raw = doc.get("sheet_header_slugs")
    if isinstance(raw, dict):
        return raw
    doc["sheet_header_slugs"] = {}
    return doc["sheet_header_slugs"]


def _load_sheet_header_slug_map(doc: dict[str, Any]) -> dict[str, str]:
    """把 YAML sheet_header_slugs 载入为 norm(header) -> slug。"""
    raw = _ensure_sheet_header_slugs(doc)
    out: dict[str, str] = {}
    for k, v in raw.items():
        if not isinstance(k, str) or not isinstance(v, str):
            continue
        ks, vs = k.strip(), v.strip()
        if not ks or not vs:
            continue
        out[_norm(ks)] = vs
    return out


def _get_or_create_slug_for_header(
    *,
    header: str,
    slug_map_norm: dict[str, str],
    yaml_sheet_header_slugs: dict[str, Any],
) -> str:
    """
    由 YAML sheet_header_slugs 决定字段 key；若缺失则自动生成并写回 YAML。
    自动生成优先使用 SUGGESTED_SHEET_HEADER_SLUGS，其次 suggest_field_slug（拉丁化优先，pypinyin 兜底）。
    """
    hn = _norm(header)
    taken = set(slug_map_norm.values())
    if not hn:
        return suggest_field_slug(header, taken_slugs=taken)
    if hn in slug_map_norm:
        return slug_map_norm[hn]

    suggested = SUGGESTED_SHEET_HEADER_SLUGS.get(hn)
    slug = (
        suggested.strip()
        if isinstance(suggested, str) and suggested.strip()
        else suggest_field_slug(header, taken_slugs=taken)
    )
    # 写回 YAML：用“首次遇到的原始表头文本”作为 key，便于用户直接肉眼维护
    yaml_sheet_header_slugs[str(header).strip()] = slug
    slug_map_norm[hn] = slug
    return slug


def _parse_sheet_block_aliases(doc: dict[str, Any], sk: str) -> dict[str, list[str]]:
    """当前 YAML 中该 sheet 下全部字段的别名列表（含空列表），用于在 default 之外的二次匹配。"""
    out: dict[str, list[str]] = {}
    sheets_raw = doc.get("sheets")
    if not isinstance(sheets_raw, dict):
        return out
    eb = sheets_raw.get(sk)
    if not isinstance(eb, dict):
        return out
    for k, v in eb.items():
        if not isinstance(k, str) or not isinstance(v, list):
            continue
        key = k.strip()
        if not key:
            continue
        out[key] = [str(x).strip() for x in v if x is not None and str(x).strip()]
    return out


def _match_header_to_existing_aliases(header: str, existing: dict[str, list[str]]) -> str | None:
    hn = _norm(header)
    if not hn:
        return None
    for field in sorted(existing.keys()):
        for a in existing.get(field, []):
            if _norm(a) == hn:
                return field
    return None


def _uniq_sorted(items: list[str]) -> list[str]:
    return sorted(set(items), key=lambda x: (len(x), x))


def _iter_sample_xlsx_paths(data_dirs: list[Path]) -> Iterator[Path]:
    """按目录顺序、每目录内路径排序遍历样例 xlsx（跳过 Excel 临时文件 ~$*.xlsx）。"""
    for data_dir in data_dirs:
        for path in sorted(data_dir.rglob("*.xlsx")):
            if path.name.startswith("~$"):
                continue
            yield path


def _parse_existing_nonempty(doc: dict[str, Any], sk: str) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    sheets_raw = doc.get("sheets")
    if not isinstance(sheets_raw, dict):
        return out
    eb = sheets_raw.get(sk)
    if not isinstance(eb, dict):
        return out
    for k, v in eb.items():
        if not isinstance(k, str) or not isinstance(v, list):
            continue
        cleaned = [str(x).strip() for x in v if x is not None and str(x).strip()]
        if cleaned:
            out[k.strip()] = cleaned
    return out


def _scan_sample(
    data_dirs: list[Path],
    sk_target: str,
    match: str,
    field_order: list[str],
    default_aliases: dict[str, list[str]],
    existing_sheet_aliases: dict[str, list[str]],
    slug_map_norm: dict[str, str],
    yaml_sheet_header_slugs: dict[str, Any],
    *,
    max_excel_files: int = 0,
) -> tuple[int, dict[str, set[str]], list[str]]:
    per_field_obs: dict[str, set[str]] = defaultdict(set)
    auto_slugs: list[str] = []
    sheet_hits = 0
    for idx, path in enumerate(_iter_sample_xlsx_paths(data_dirs)):
        if max_excel_files and idx >= max_excel_files:
            break
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        try:
            for sn in wb.sheetnames:
                n = _norm(sn)
                if match == "exact":
                    ok = n == _norm(sk_target)
                else:
                    ok = _norm(sk_target) in n
                if not ok:
                    continue
                sheet_hits += 1
                ws = wb[sn]
                hdr = _detect_header_row(ws)
                if not hdr:
                    continue
                for h in hdr:
                    t = str(h).strip()
                    if not t:
                        continue
                    f = _header_to_field(t, field_order, default_aliases)
                    if f is None:
                        f = _match_header_to_existing_aliases(t, existing_sheet_aliases)
                    if f is None:
                        slug = _get_or_create_slug_for_header(
                            header=t,
                            slug_map_norm=slug_map_norm,
                            yaml_sheet_header_slugs=yaml_sheet_header_slugs,
                        )
                        per_field_obs[slug].add(t)
                        auto_slugs.append(slug)
                    else:
                        per_field_obs[f].add(t)
        finally:
            wb.close()
    return sheet_hits, per_field_obs, sorted(set(auto_slugs))


def _sheet_name_matches_key(sn: str, sk: str, match: str) -> bool:
    n = _norm(sn)
    if match == "exact":
        return n == _norm(sk)
    return _norm(sk) in n


def _scan_all_sheet_keys_one_pass(
    data_dirs: list[Path],
    keys: list[str],
    match: str,
    field_order: list[str],
    default_aliases: dict[str, list[str]],
    doc: dict[str, Any],
    slug_map_norm: dict[str, str],
    yaml_sheet_header_slugs: dict[str, Any],
    *,
    max_excel_files: int = 0,
    progress_every: int = 100,
) -> tuple[dict[str, int], dict[str, dict[str, set[str]]], dict[str, list[str]]]:
    """
    对所有 sheet key 只遍历样例目录一次（每个 xlsx 只打开一次），汇总各 key 的列观测。
    返回：hits_per_sk, per_field_obs_per_sk, auto_slugs_per_sk（去重后的 slug 列表）。
    """
    existing_by_sk: dict[str, dict[str, list[str]]] = {
        sk: _parse_sheet_block_aliases(doc, sk) for sk in keys
    }
    hits: dict[str, int] = {sk: 0 for sk in keys}
    obs: dict[str, dict[str, set[str]]] = {sk: defaultdict(set) for sk in keys}
    auto_lists: dict[str, list[str]] = {sk: [] for sk in keys}

    last_index = -1
    for idx, path in enumerate(_iter_sample_xlsx_paths(data_dirs)):
        if max_excel_files and idx >= max_excel_files:
            break
        last_index = idx
        if progress_every and idx > 0 and idx % progress_every == 0:
            print(f"  ... scanned {idx} xlsx files", flush=True)
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        try:
            for sn in wb.sheetnames:
                matched = [sk for sk in keys if _sheet_name_matches_key(sn, sk, match)]
                if not matched:
                    continue
                ws = wb[sn]
                hdr = _detect_header_row(ws)
                if not hdr:
                    continue
                for sk in matched:
                    hits[sk] += 1
                    ex = existing_by_sk[sk]
                    for h in hdr:
                        t = str(h).strip()
                        if not t:
                            continue
                        f = _header_to_field(t, field_order, default_aliases)
                        if f is None:
                            f = _match_header_to_existing_aliases(t, ex)
                        if f is None:
                            slug = _get_or_create_slug_for_header(
                                header=t,
                                slug_map_norm=slug_map_norm,
                                yaml_sheet_header_slugs=yaml_sheet_header_slugs,
                            )
                            obs[sk][slug].add(t)
                            auto_lists[sk].append(slug)
                        else:
                            obs[sk][f].add(t)
        finally:
            wb.close()

    auto_slugs_out = {sk: sorted(set(auto_lists[sk])) for sk in keys}
    print(f"  done, scanned {last_index + 1} xlsx file(s)", flush=True)
    return hits, obs, auto_slugs_out


def _build_block(
    *,
    field_order: list[str],
    default_aliases: dict[str, list[str]],
    per_field_obs: dict[str, set[str]],
    existing_sheet: dict[str, list[str]],
    emit_all_default_keys: bool,
    core_fields: list[str],
    has_any_observation: bool,
) -> dict[str, list[str]]:
    new_block: dict[str, list[str]] = {}

    field_order_set = set(field_order)

    if emit_all_default_keys:
        for f in field_order:
            dset = set(default_aliases.get(f, []))
            oset = per_field_obs.get(f, set()) or set()
            merged = dset | oset
            if merged == dset:
                new_block[f] = []
            else:
                new_block[f] = _uniq_sorted(list(merged))
        for f, oset in sorted(per_field_obs.items(), key=lambda x: x[0]):
            if f in field_order_set or not oset:
                continue
            dset = set(default_aliases.get(f, []))
            merged = dset | oset
            if merged == dset:
                new_block[f] = []
            else:
                new_block[f] = _uniq_sorted(list(merged))
        return new_block

    if not has_any_observation:
        # 无样例：通用 9 字段 + 保留已有非空手工项（可含额外字段 key）
        for f in core_fields:
            new_block[f] = []
        for f, als in existing_sheet.items():
            if als:
                new_block[f] = list(als)
        return new_block

    for f in field_order:
        dset = set(default_aliases.get(f, []))
        oset = per_field_obs.get(f, set()) or set()

        if not oset:
            if f in existing_sheet:
                new_block[f] = list(existing_sheet[f])
            continue

        merged = dset | oset
        if merged == dset:
            new_block[f] = []
        else:
            new_block[f] = _uniq_sorted(list(merged))

    for f, oset in sorted(per_field_obs.items(), key=lambda x: x[0]):
        if f in field_order_set or not oset:
            continue
        dset = set(default_aliases.get(f, []))
        merged = dset | oset
        if merged == dset:
            new_block[f] = []
        else:
            new_block[f] = _uniq_sorted(list(merged))

    return new_block


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--data",
        action="append",
        default=None,
        help="样例根目录（可重复传入以合并多个目录；缺省为项目内 Source_Data/Invoice）",
    )
    ap.add_argument("--sheet-key", type=str, default="发票基础信息", help="写入 YAML 的 sheets 下的 key（与 --all-sheet-keys 互斥）")
    ap.add_argument(
        "--all-sheet-keys",
        action="store_true",
        help="按 config/sheet_mapping.yaml 中全部 sheet key 依次处理并写回 sheets",
    )
    ap.add_argument(
        "--match",
        choices=["contains", "exact"],
        default="contains",
        help="工作表名匹配：contains=归一后包含 sheet-key；exact=归一后等于 sheet-key",
    )
    ap.add_argument("--config", type=str, default=str(CFG_PATH), help="field_mapping.yaml 路径")
    ap.add_argument(
        "--sheet-mapping",
        type=str,
        default=str(SHEET_MAP_PATH),
        help="sheet_mapping.yaml 路径（--all-sheet-keys 时使用）",
    )
    ap.add_argument(
        "--emit-all-default-keys",
        action="store_true",
        help="把 default 中所有字段 key 都写入该 sheet 块（未观测到的写 []）",
    )
    ap.add_argument(
        "--core-fields",
        type=str,
        default="",
        help="无样例时的通用字段 key，逗号分隔；留空则使用内置 9 个核心字段",
    )
    ap.add_argument(
        "--max-excel-files",
        type=int,
        default=0,
        help="最多扫描多少个 xlsx（按路径顺序，跨 --data 累计）；0 表示不限制",
    )
    args = ap.parse_args()

    data_roots = args.data if args.data else [str(DATA_DEFAULT)]
    data_dirs: list[Path] = []
    for r in data_roots:
        p = Path(r)
        if not p.is_absolute():
            p = ROOT / p
        if not p.is_dir():
            raise SystemExit(f"not a directory: {p}")
        data_dirs.append(p)
    cfg_path = Path(args.config)
    if not cfg_path.is_absolute():
        cfg_path = ROOT / cfg_path
    sm_path = Path(args.sheet_mapping)
    if not sm_path.is_absolute():
        sm_path = ROOT / sm_path

    core_fields = DEFAULT_CORE_FIELDS
    if str(args.core_fields).strip():
        core_fields = [x.strip() for x in str(args.core_fields).split(",") if x.strip()]

    raw = cfg_path.read_text(encoding="utf-8")
    doc = yaml.safe_load(raw)
    if not isinstance(doc, dict):
        raise SystemExit("invalid yaml root")
    field_order, default_aliases = _default_field_order_and_aliases(doc)
    if not field_order:
        raise SystemExit("no default: block")

    if "sheets" not in doc or not isinstance(doc["sheets"], dict):
        doc["sheets"] = {}

    yaml_sheet_header_slugs = _ensure_sheet_header_slugs(doc)
    slug_map_norm = _load_sheet_header_slug_map(doc)

    if args.all_sheet_keys:
        keys = _load_sheet_mapping_keys(sm_path)
        if not keys:
            raise SystemExit(f"no keys in {sm_path}")
        print(
            f"Scanning {len(data_dirs)} data root(s), max_excel_files={args.max_excel_files or 'unlimited'} ...",
            flush=True,
        )
        hits_all, obs_all, auto_all = _scan_all_sheet_keys_one_pass(
            data_dirs,
            keys,
            args.match,
            field_order,
            default_aliases,
            doc,
            slug_map_norm,
            yaml_sheet_header_slugs,
            max_excel_files=args.max_excel_files,
        )
        new_sheets: dict[str, Any] = {}
        summary: list[str] = []
        for sk in keys:
            existing_sheet = _parse_existing_nonempty(doc, sk)
            per_field_obs = obs_all.get(sk) or {}
            auto_slugs = auto_all.get(sk) or []
            hits = hits_all.get(sk, 0)
            has_obs = any(bool(s) for s in per_field_obs.values())
            block = _build_block(
                field_order=field_order,
                default_aliases=default_aliases,
                per_field_obs=dict(per_field_obs),
                existing_sheet=existing_sheet,
                emit_all_default_keys=args.emit_all_default_keys,
                core_fields=core_fields,
                has_any_observation=has_obs,
            )
            new_sheets[sk] = block
            mode = "sample" if has_obs else "generic"
            extra = f" sheet_only_fields={len(auto_slugs)}" if auto_slugs else ""
            summary.append(f"{sk}: hits={hits} mode={mode} keys={len(block)}{extra}")
        # 保留 mapping 外的 sheet 块（若有）
        for k, v in doc["sheets"].items():
            if k not in new_sheets and isinstance(v, dict):
                new_sheets[k] = v
        doc["sheets"] = new_sheets
    else:
        sk_target = args.sheet_key.strip()
        existing_sheet = _parse_existing_nonempty(doc, sk_target)
        existing_aliases = _parse_sheet_block_aliases(doc, sk_target)
        hits, per_field_obs, auto_slugs = _scan_sample(
            data_dirs,
            sk_target,
            args.match,
            field_order,
            default_aliases,
            existing_aliases,
            slug_map_norm,
            yaml_sheet_header_slugs,
            max_excel_files=args.max_excel_files,
        )
        has_obs = any(bool(s) for s in per_field_obs.values())
        new_block = _build_block(
            field_order=field_order,
            default_aliases=default_aliases,
            per_field_obs=dict(per_field_obs),
            existing_sheet=existing_sheet,
            emit_all_default_keys=args.emit_all_default_keys,
            core_fields=core_fields,
            has_any_observation=has_obs,
        )
        doc["sheets"][sk_target] = new_block
        suf = f" sheet_only_fields={auto_slugs}" if auto_slugs else ""
        summary = [
            f"{sk_target}: hits={hits} mode={'sample' if has_obs else 'generic'} keys={len(new_block)}{suf}"
        ]
        _auto_slugs_single = auto_slugs

    header = ""
    lines = raw.splitlines()
    if lines and lines[0].startswith("#"):
        header = lines[0] + "\n"

    out = header + yaml.safe_dump(doc, allow_unicode=True, sort_keys=False, default_flow_style=False)
    cfg_path.write_text(out, encoding="utf-8")

    for line in summary:
        print(line)
    if not args.all_sheet_keys and _auto_slugs_single:
        print("  sheet-only field keys (default 未声明，由样例/slug 生成):", ", ".join(_auto_slugs_single))
    print(f"written {cfg_path}")


if __name__ == "__main__":
    main()
